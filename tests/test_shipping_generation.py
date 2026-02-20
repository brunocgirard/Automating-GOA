"""Shipping generation behavior tests."""

from __future__ import annotations

import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document

from api.services import shipping_doc_service as shipping_service


def _sample_shipping_state() -> dict[str, Any]:
    return {
        "quoteId": 42,
        "quoteRef": "Q-42",
        "client": {},
        "machines": [
            {
                "id": "machine-1",
                "machineId": 1,
                "lineItemOptionId": None,
                "machineName": "Machine One",
                "model": "Machine One",
                "hsCode": "",
                "serialNumber": "",
                "unitPrice": 1000,
                "truckId": "truck-1",
                "crates": [],
            },
            {
                "id": "machine-2",
                "machineId": 2,
                "lineItemOptionId": None,
                "machineName": "Machine Two",
                "model": "Machine Two",
                "hsCode": "",
                "serialNumber": "",
                "unitPrice": 2000,
                "truckId": "truck-2",
                "crates": [],
            },
        ],
        "trucks": [
            {"id": "truck-1", "name": "Truck 1"},
            {"id": "truck-2", "name": "Truck 2"},
        ],
        "meta": {},
    }


def _single_truck_shipping_state() -> dict[str, Any]:
    payload = _sample_shipping_state()
    payload["machines"] = [
        machine
        for machine in payload["machines"]
        if machine.get("truckId") == "truck-1"
    ]
    payload["trucks"] = [{"id": "truck-1", "name": "Truck 1"}]
    return payload


def _fake_generate_single_docx(_doc_type: str, truck_state: dict[str, Any], output_path: Path) -> None:
    machine_count = len([m for m in truck_state.get("machines", []) if isinstance(m, dict)])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(f"machines={machine_count}", encoding="utf-8")


def test_generate_all_documents_creates_set_per_truck(tmp_path, monkeypatch):
    monkeypatch.setattr(shipping_service, "OUTPUT_ROOT", tmp_path)
    monkeypatch.setattr(shipping_service, "_generate_single_docx", _fake_generate_single_docx)

    artifact = shipping_service.generate_shipping_documents(
        shipping_data=_sample_shipping_state(),
        quote_ref="Q-42",
        document_type="all",
    )

    assert artifact.media_type == "application/zip"
    assert artifact.filename.endswith("_shipping_documents.zip")

    with zipfile.ZipFile(artifact.path, "r") as archive:
        names = sorted(archive.namelist())
        assert names == sorted(
            [
                "packing-slip-truck-1.docx",
                "commercial-invoice-truck-1.docx",
                "certificate-of-origin-truck-1.docx",
                "packing-slip-truck-2.docx",
                "commercial-invoice-truck-2.docx",
                "certificate-of-origin-truck-2.docx",
            ]
        )


def test_generate_single_document_creates_zip_for_multiple_trucks(tmp_path, monkeypatch):
    monkeypatch.setattr(shipping_service, "OUTPUT_ROOT", tmp_path)
    monkeypatch.setattr(shipping_service, "_generate_single_docx", _fake_generate_single_docx)

    artifact = shipping_service.generate_shipping_documents(
        shipping_data=_sample_shipping_state(),
        quote_ref="Q-42",
        document_type="packing_slip",
    )

    assert artifact.media_type == "application/zip"
    assert artifact.filename.endswith("_packing_slip_by_truck.zip")

    with zipfile.ZipFile(artifact.path, "r") as archive:
        names = sorted(archive.namelist())
        assert names == ["packing-slip-truck-1.docx", "packing-slip-truck-2.docx"]


def test_generate_single_html_document_for_one_truck(tmp_path, monkeypatch):
    monkeypatch.setattr(shipping_service, "OUTPUT_ROOT", tmp_path)

    artifact = shipping_service.generate_shipping_documents(
        shipping_data=_single_truck_shipping_state(),
        quote_ref="Q-42",
        document_type="packing_slip",
        output_format="html",
    )

    assert artifact.media_type.startswith("text/html")
    assert artifact.filename.endswith(".html")
    content = artifact.path.read_text(encoding="utf-8")
    assert "<!doctype html>" in content.lower()
    assert "Packing Slip" in content


def test_generate_all_html_documents_creates_set_per_truck(tmp_path, monkeypatch):
    monkeypatch.setattr(shipping_service, "OUTPUT_ROOT", tmp_path)

    artifact = shipping_service.generate_shipping_documents(
        shipping_data=_sample_shipping_state(),
        quote_ref="Q-42",
        document_type="all",
        output_format="html",
    )

    assert artifact.media_type == "application/zip"
    assert artifact.filename.endswith("_shipping_documents.zip")

    with zipfile.ZipFile(artifact.path, "r") as archive:
        names = sorted(archive.namelist())
        assert names == sorted(
            [
                "packing-slip-truck-1.html",
                "commercial-invoice-truck-1.html",
                "certificate-of-origin-truck-1.html",
                "packing-slip-truck-2.html",
                "commercial-invoice-truck-2.html",
                "certificate-of-origin-truck-2.html",
            ]
        )


def test_replace_tokens_in_paragraph_supports_double_brace_placeholders():
    doc = Document()
    paragraph = doc.add_paragraph("PO {{Customer_PO}} / Name {{Customer}} / OX {{Ox}}")

    shipping_service._replace_tokens_in_paragraph(
        paragraph,
        {"Customer_PO": "PO-100", "Customer": "ACME", "Ox": "OX-55"},
        {},
    )

    assert paragraph.text == "PO PO-100 / Name ACME / OX OX-55"


def test_fill_repeating_machine_rows_without_next_record_marker():
    doc = Document()
    table = doc.add_table(rows=2, cols=1)
    table.rows[0].cells[0].text = "Header"
    table.rows[1].cells[0].text = "Machine {{Machine}} / Serial {{Serial_Number}} / HS {{HS}} / Qty {{Qty}}"

    machine_rows = [
        {"Machine": "M-1", "Serial_Number": "S-1", "HS": "H-1", "Qty": "1", "TYPE": "", "Next Record": ""},
        {"Machine": "M-2", "Serial_Number": "S-2", "HS": "H-2", "Qty": "1", "TYPE": "", "Next Record": ""},
    ]

    shipping_service._fill_repeating_machine_rows(doc, machine_rows)

    assert len(table.rows) == 3
    assert table.rows[1].cells[0].text == "Machine M-1 / Serial S-1 / HS H-1 / Qty 1"
    assert table.rows[2].cells[0].text == "Machine M-2 / Serial S-2 / HS H-2 / Qty 1"


def test_replace_tokens_in_document_replaces_header_placeholders():
    doc = Document()
    header = doc.sections[0].header
    header.paragraphs[0].text = "Date: {{Date}} / Year: {{Year}} / OX: {{Ox}}"

    shipping_service._replace_tokens_in_document(
        doc,
        {"Date": "2026-02-20", "Year": "2026", "Ox": "OX-123"},
        {},
    )

    assert header.paragraphs[0].text == "Date: 2026-02-20 / Year: 2026 / OX: OX-123"


def test_resolve_template_path_prefers_paking_slip1_when_present(monkeypatch, tmp_path):
    preferred = tmp_path / "Paking Slip1.docx"
    fallback = tmp_path / "Paking Slip.docx"
    preferred.write_bytes(b"preferred")
    fallback.write_bytes(b"fallback")

    monkeypatch.setattr(shipping_service, "PACKING_SLIP_TEMPLATE_CANDIDATES", (preferred, fallback))
    monkeypatch.setitem(shipping_service.TEMPLATE_PATHS, "packing_slip", fallback)

    assert shipping_service._resolve_template_path("packing_slip") == preferred


def test_resolve_template_path_prefers_commercial_invoice1_when_present(monkeypatch, tmp_path):
    preferred = tmp_path / "Commercial Invoice1.docx"
    fallback = tmp_path / "Commercial Invoice.docx"
    preferred.write_bytes(b"preferred")
    fallback.write_bytes(b"fallback")

    monkeypatch.setattr(shipping_service, "COMMERCIAL_INVOICE_TEMPLATE_CANDIDATES", (preferred, fallback))
    monkeypatch.setitem(shipping_service.TEMPLATE_PATHS, "commercial_invoice", fallback)

    assert shipping_service._resolve_template_path("commercial_invoice") == preferred


def test_resolve_template_path_prefers_certificate_origin1_when_present(monkeypatch, tmp_path):
    preferred = tmp_path / "CERTIFICATION OF ORIGIN_NAFTA1.docx"
    fallback = tmp_path / "CERTIFICATION OF ORIGIN_NAFTA.docx"
    preferred.write_bytes(b"preferred")
    fallback.write_bytes(b"fallback")

    monkeypatch.setattr(shipping_service, "CERTIFICATE_ORIGIN_TEMPLATE_CANDIDATES", (preferred, fallback))
    monkeypatch.setitem(shipping_service.TEMPLATE_PATHS, "certificate_origin", fallback)

    assert shipping_service._resolve_template_path("certificate_origin") == preferred


def test_build_global_token_replacements_uses_current_date_for_date_and_year():
    now = datetime.now()
    shipping_data = {
        "client": {
            "customerName": "ACME",
            "company": "ACME",
            "orderDate": "2001-01-01",
        },
        "meta": {},
    }
    machine_rows = [{"Machine": "M-1", "Serial_Number": "S-1", "HS": "H-1", "Qty": "1", "TYPE": ""}]

    replacements = shipping_service._build_global_token_replacements(shipping_data, machine_rows)

    assert replacements["Date"] == now.strftime("%Y-%m-%d")
    assert replacements["Year"] == str(now.year)
    try:
        expected_next_year = now.replace(year=now.year + 1).strftime("%Y-%m-%d")
    except ValueError:
        expected_next_year = now.replace(month=2, day=28, year=now.year + 1).strftime("%Y-%m-%d")
    assert replacements["Next_Year_Date"] == expected_next_year
    # Keep orderDate mapping for templates that intentionally use that token.
    assert replacements["Order_date"] == "2001-01-01"


def test_machine_rows_include_price_total_price_and_lowercase_type():
    shipping_data = {
        "machines": [
            {
                "model": "Inline Filler",
                "serialNumber": "SN-1",
                "hsCode": "8422.30",
                "type": "Bottle",
                "qty": "2",
                "unitPrice": 1250.5,
            }
        ]
    }

    rows = shipping_service._machine_rows(shipping_data)

    assert rows[0]["type"] == "Bottle"
    assert rows[0]["Price"] == "$1,250.50"
    assert rows[0]["Total_Price"] == "$2,501.00"


def test_build_global_token_replacements_maps_total_price_and_type_tokens():
    shipping_data = {
        "client": {
            "customerName": "ACME",
            "orderDate": "2001-01-01",
        },
        "machines": [{"unitPrice": 1000.0}, {"unitPrice": 2000.0}],
        "meta": {"totalInvoiceAmount": "3500"},
    }
    machine_rows = [
        {
            "Machine": "M-1",
            "Serial_Number": "S-1",
            "HS": "H-1",
            "Qty": "2",
            "TYPE": "Bottle",
            "type": "Bottle",
            "Price": "$1,000.00",
            "Total_Price": "$2,000.00",
        }
    ]

    replacements = shipping_service._build_global_token_replacements(shipping_data, machine_rows)

    assert replacements["type"] == "Bottle"
    assert replacements["Price"] == "$1,000.00"
    assert replacements["Total_Price"] == "$3,500.00"
