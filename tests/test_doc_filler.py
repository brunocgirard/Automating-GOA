import logging
from docx import Document
from src.utils.doc_filler import fill_word_document_from_llm_data


def create_simple_template(path):
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "{{existing}}"
    doc.save(path)


def test_warning_for_missing_placeholder(tmp_path, caplog):
    template = tmp_path / "template.docx"
    create_simple_template(template)
    output = tmp_path / "out.docx"

    with caplog.at_level(logging.WARNING):
        fill_word_document_from_llm_data(str(template), {"missing": "value"}, str(output))

    warnings = [record for record in caplog.records if record.levelno == logging.WARNING]
    assert any("missing" in w.getMessage() and "table 0" in w.getMessage() for w in warnings)

