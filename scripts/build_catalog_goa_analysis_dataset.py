from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from src.utils.pdf_utils import extract_full_pdf_text, extract_line_item_details


def _normalize_text(text: str) -> str:
    return " ".join(str(text or "").split()).strip()


def _extract_goa_content(docx_path: Path) -> dict[str, Any]:
    doc = Document(str(docx_path))

    paragraph_lines = [_normalize_text(p.text) for p in doc.paragraphs]
    paragraph_lines = [line for line in paragraph_lines if line]

    table_rows: list[dict[str, Any]] = []
    flattened_lines: list[str] = []

    for table_index, table in enumerate(doc.tables):
        for row_index, row in enumerate(table.rows):
            cells = [_normalize_text(cell.text) for cell in row.cells]
            if not any(cells):
                continue
            table_rows.append(
                {
                    "table_index": table_index,
                    "row_index": row_index,
                    "cells": cells,
                }
            )
            flattened_lines.append(" | ".join(cell for cell in cells if cell))

    # Keep a single searchable text block for simple phrase scans later.
    flattened_text = "\n".join([*paragraph_lines, *flattened_lines]).strip()

    return {
        "paragraph_count": len(paragraph_lines),
        "table_count": len(doc.tables),
        "table_row_count": len(table_rows),
        "paragraph_lines": paragraph_lines,
        "table_rows": table_rows,
        "flattened_text": flattened_text,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Build an analysis dataset by combining the catalog manifest with extracted "
            "quote PDF text and flattened completed GOA content."
        )
    )
    parser.add_argument(
        "--manifest",
        default="output/catalog_goa_manifest.json",
        help="Path to the catalog manifest JSON produced by build_catalog_goa_manifest.py",
    )
    parser.add_argument(
        "--output",
        default="output/catalog_goa_analysis_dataset.json",
        help="Path to write the enriched analysis dataset JSON.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    manifest_path = Path(args.manifest).resolve()
    if not manifest_path.exists():
        raise FileNotFoundError(f"Manifest not found: {manifest_path}")

    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    source_dir = Path(manifest["source_dir"])
    if not source_dir.exists():
        raise FileNotFoundError(f"Source directory from manifest not found: {source_dir}")

    pdf_cache: dict[str, dict[str, Any]] = {}
    enriched_groups: list[dict[str, Any]] = []
    missing_files: list[dict[str, str]] = []
    total_pdf_chars = 0
    total_line_items = 0
    total_goa_chars = 0
    total_goa_rows = 0

    for group in manifest.get("quote_groups", []):
        pdf_documents: list[dict[str, Any]] = []
        for relative_pdf_path in group.get("pdf_files", []):
            cached = pdf_cache.get(relative_pdf_path)
            if cached is None:
                pdf_path = source_dir / relative_pdf_path
                if not pdf_path.exists():
                    missing_files.append({"type": "pdf", "relative_path": relative_pdf_path})
                    continue
                full_text = extract_full_pdf_text(str(pdf_path))
                line_items = extract_line_item_details(str(pdf_path))
                line_item_descriptions = [
                    _normalize_text(item.get("description") or "")
                    for item in line_items
                    if _normalize_text(item.get("description") or "")
                ]
                cached = {
                    "relative_path": relative_pdf_path,
                    "text_char_count": len(full_text),
                    "page_text": full_text,
                    "line_item_count": len(line_items),
                    "line_item_descriptions": line_item_descriptions,
                    "line_items": line_items,
                }
                pdf_cache[relative_pdf_path] = cached
                total_pdf_chars += len(full_text)
                total_line_items += len(line_items)
            pdf_documents.append(cached)

        records: list[dict[str, Any]] = []
        for record in group.get("records", []):
            relative_goa_path = record["relative_path"]
            goa_path = source_dir / relative_goa_path
            if not goa_path.exists():
                missing_files.append({"type": "goa", "relative_path": relative_goa_path})
                continue
            goa_content = _extract_goa_content(goa_path)
            total_goa_chars += len(goa_content["flattened_text"])
            total_goa_rows += int(goa_content["table_row_count"])

            enriched_record = dict(record)
            enriched_record["goa_content"] = goa_content
            records.append(enriched_record)

        enriched_groups.append(
            {
                "quote_key": group["quote_key"],
                "pdf_files": group.get("pdf_files", []),
                "goa_files": group.get("goa_files", []),
                "record_count": group.get("record_count", 0),
                "has_multiple_goas": group.get("has_multiple_goas", False),
                "pdf_documents": pdf_documents,
                "records": records,
            }
        )

    dataset = {
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "manifest_path": str(manifest_path),
        "source_dir": str(source_dir),
        "totals": {
            "quote_group_count": len(enriched_groups),
            "goa_record_count": sum(len(group["records"]) for group in enriched_groups),
            "pdf_document_count": sum(len(group["pdf_documents"]) for group in enriched_groups),
            "total_pdf_text_chars": total_pdf_chars,
            "total_line_item_count": total_line_items,
            "total_goa_text_chars": total_goa_chars,
            "total_goa_table_rows": total_goa_rows,
            "missing_file_count": len(missing_files),
        },
        "missing_files": missing_files,
        "quote_groups": enriched_groups,
    }

    output_path = Path(args.output).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(dataset, indent=2), encoding="utf-8")

    totals = dataset["totals"]
    print(f"Wrote analysis dataset to {output_path}")
    print(
        "Summary: "
        f"quote_groups={totals['quote_group_count']} "
        f"goa_records={totals['goa_record_count']} "
        f"pdfs={totals['pdf_document_count']} "
        f"pdf_chars={totals['total_pdf_text_chars']} "
        f"line_items={totals['total_line_item_count']} "
        f"goa_chars={totals['total_goa_text_chars']} "
        f"missing={totals['missing_file_count']}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
