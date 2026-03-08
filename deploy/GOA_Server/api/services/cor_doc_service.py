"""COR workflow helpers for prefill and DOCX generation."""

from __future__ import annotations

import re
from copy import deepcopy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt

from api.services._doc_helpers import (
    current_date as _current_date,
    replace_tokens_in_container,
    slugify as _slugify,
    to_float as _to_float,
    to_text as _to_text,
)

OUTPUT_ROOT = Path("data/generated/cor")
TEMPLATE_PATH = Path("Mail_merge/Customer_Oxxxx_COR.docx")
TEMPLATE_PATH_CANDIDATES = (
    Path("Mail_merge/Customer_Oxxxx_COR1.docx"),
    Path("Mail_merge/Customer_Oxxxx_COR.docx"),
)
INSERTED_TEXT_FONT_SIZE_PT = 8


@dataclass
class GeneratedCorArtifact:
    path: Path
    filename: str
    media_type: str


def build_cor_prefill_data(
    quote: dict[str, Any],
) -> dict[str, Any]:
    """Build a prefilled COR state from quote info only.

    COR line rows must always be user-entered because this document
    is used to add/remove items, not mirror quote priced items.
    """
    line_items = [
        {
            "id": "cor-line-1",
            "qty": "",
            "reqDescription": "",
            "unitCost": "",
            "selectedItems": "",
        }
    ]

    return {
        "quoteId": int(quote.get("id") or 0),
        "quoteRef": _to_text(quote.get("quote_ref")),
        "client": {
            "company": _to_text(quote.get("customer_name")) or _to_text(quote.get("company")),
            "customerPO": _to_text(quote.get("customer_po")),
            "orderDate": _to_text(quote.get("order_date")),
            "ax": _to_text(quote.get("ax")),
            "ox": _to_text(quote.get("ox")),
            "machine": _to_text(quote.get("machine_model")),
        },
        "corNo": "1",
        "revisionDescription": "",
        "corStatus": "",
        "capmaticPM": "",
        "initiatorOfChange": "contact_person",
        "salesRep": "",
        "contactPerson": _to_text(quote.get("company")),
        "impactDeliverables": "",
        "impactDeliveryDate": "",
        "paymentTerms": "",
        "currency": "",
        "approvalDate": _current_date(),
        "justificationForChange": "",
        "comments": "",
        "lineItems": line_items,
    }


def generate_cor_document(cor_data: dict[str, Any], quote_ref: str) -> GeneratedCorArtifact:
    """Generate COR DOCX document from COR state."""
    template_path = _resolve_template_path()
    if not template_path.exists():
        candidates = ", ".join(str(path) for path in TEMPLATE_PATH_CANDIDATES)
        raise FileNotFoundError(f"Template not found. Checked: {candidates}")

    safe_quote_ref = _slugify(quote_ref) or "quote"
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_dir = OUTPUT_ROOT / safe_quote_ref / timestamp
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"{safe_quote_ref}_cor.docx"

    document = Document(str(template_path))
    _replace_core_tokens(document, cor_data)
    _fill_cor_template_tables(document, cor_data)
    document.save(str(output_path))

    return GeneratedCorArtifact(
        path=output_path,
        filename=output_path.name,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


def _resolve_template_path() -> Path:
    for candidate in TEMPLATE_PATH_CANDIDATES:
        if candidate.exists():
            return candidate
    return TEMPLATE_PATH


def _replace_core_tokens(document: Document, cor_data: dict[str, Any]) -> None:
    client = cor_data.get("client") if isinstance(cor_data.get("client"), dict) else {}
    cor_no = _to_text(cor_data.get("corNo"))
    comments = _to_text(cor_data.get("comments"))
    cor_status = _to_text(cor_data.get("corStatus"))
    capmatic_pm = _to_text(cor_data.get("capmaticPM"))
    initiator_of_change = _to_text(cor_data.get("initiatorOfChange"))
    sales_rep = _to_text(cor_data.get("salesRep"))
    contact_person = _to_text(cor_data.get("contactPerson"))
    initiator_value = _resolve_initiator_value(contact_person, capmatic_pm, initiator_of_change)
    order_date = _to_text(client.get("orderDate"))
    payment_terms = _to_text(cor_data.get("paymentTerms"))
    currency = _to_text(cor_data.get("currency"))
    approval_date = _current_date()
    justification = _to_text(cor_data.get("justificationForChange"))
    cor_no_label = _format_cor_no_label(cor_no)
    replacements = {
        "Customer": _to_text(client.get("company")),
        "Company": _to_text(client.get("company")),
        "Customer_PO": _to_text(client.get("customerPO")),
        "Ax": _to_text(client.get("ax")),
        "Ox": _to_text(client.get("ox")),
        "Machine": _to_text(client.get("machine")),
        "Order_date": order_date,
        "Contact_Person": contact_person,
        "Date": approval_date,
        "Justification_of_Change": justification,
        # Keep both spellings because template currently uses the "Satus" token.
        "COR_Satus": cor_status,
        "COR_Status": cor_status,
        "Cor_Status": cor_status,
        "COR NO. #": cor_no_label,
        "Capmatic_PM": capmatic_pm,
        "Sales_Rep": sales_rep,
        "Payment_Terms": payment_terms,
        "Currency": currency,
        "Comments": comments,
        "Comment": comments,
    }

    replace_tokens_in_container(
        document,
        replacements,
        preprocess_text=lambda text: _replace_initiator_pair(text, initiator_value),
        postprocess_text=_cleanup_or_phrase,
        recurse_nested_tables=False,
    )
    for section in document.sections:
        containers = (
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        )
        for container in containers:
            replace_tokens_in_container(
                container,
                replacements,
                preprocess_text=lambda text: _replace_initiator_pair(text, initiator_value),
                postprocess_text=_cleanup_or_phrase,
                recurse_nested_tables=False,
            )


def _fill_cor_template_tables(document: Document, cor_data: dict[str, Any]) -> None:
    tables = document.tables
    if len(tables) < 3:
        return

    client = cor_data.get("client") if isinstance(cor_data.get("client"), dict) else {}
    cor_no = _to_text(cor_data.get("corNo"))
    cor_no_label = _format_cor_no_label(cor_no)
    order_date = _to_text(client.get("orderDate"))
    cor_status = _to_text(cor_data.get("corStatus"))
    capmatic_pm = _to_text(cor_data.get("capmaticPM"))
    initiator_of_change = _to_text(cor_data.get("initiatorOfChange"))
    sales_rep = _to_text(cor_data.get("salesRep"))
    contact_person = _to_text(cor_data.get("contactPerson"))
    initiator_text = _resolve_initiator_value(contact_person, capmatic_pm, initiator_of_change)
    impact_deliverables = _to_text(cor_data.get("impactDeliverables"))
    impact_delivery_date = _to_text(cor_data.get("impactDeliveryDate"))
    payment_terms = _to_text(cor_data.get("paymentTerms"))
    currency = _to_text(cor_data.get("currency"))
    approval_date = _current_date()
    justification = _to_text(cor_data.get("justificationForChange"))
    comments = _to_text(cor_data.get("comments"))
    line_items = [entry for entry in cor_data.get("lineItems", []) if isinstance(entry, dict)]

    # Table 0 has COR no cell.
    table0 = tables[0]
    if len(table0.rows) > 3 and len(table0.rows[3].cells) > 0:
        _set_cell_text(table0.rows[3].cells[0], cor_status)
    if len(table0.rows) > 4 and len(table0.rows[4].cells) > 1:
        current_cell_text = _to_text(table0.rows[4].cells[1].text)
        cor_no_value = cor_no_label if "COR" in current_cell_text.upper() else cor_no
        _set_cell_text(table0.rows[4].cells[1], cor_no_value)

    # Table 1 has order date.
    table1 = tables[1]
    if len(table1.rows) > 1 and len(table1.rows[1].cells) > 1:
        _set_cell_text(table1.rows[1].cells[1], order_date)
    if len(table1.rows) > 1 and len(table1.rows[1].cells) > 3:
        _set_cell_text(table1.rows[1].cells[3], capmatic_pm)
    if len(table1.rows) > 2 and len(table1.rows[2].cells) > 1:
        _set_cell_text(table1.rows[2].cells[1], sales_rep)
    if len(table1.rows) > 3:
        for idx in range(1, min(len(table1.rows[3].cells), 4)):
            _set_cell_text(table1.rows[3].cells[idx], initiator_text)

    table2 = tables[2]
    _fill_justification_rows(table2, justification)
    _fill_impact_rows(table2, impact_deliverables, impact_delivery_date)
    _fill_payment_terms_row(table2, payment_terms)
    _fill_line_item_rows(table2, line_items)
    _fill_currency_cell(table2, currency)

    if len(tables) > 3:
        table3 = tables[3]
        if len(table3.rows) > 1 and len(table3.rows[1].cells) > 1:
            _set_cell_text(table3.rows[1].cells[1], capmatic_pm)
        if len(table3.rows) > 1 and len(table3.rows[1].cells) > 3:
            _set_cell_text(table3.rows[1].cells[3], approval_date)

    _fill_comments_rows(document, comments)


def _fill_justification_rows(table: Any, justification: str) -> None:
    if len(table.rows) < 2:
        return

    # Justification starts at row 1 and ends before the first section marker row.
    marker_rows = [
        idx
        for marker in (
            "PAYMENT TERMS",
            "IMPACT OF CHANGE TO DELIVERABLES",
            "IMPACT ON DELIVERY DATE",
            "Qty. Req",
        )
        if (idx := _find_row_index_containing(table, marker)) is not None and idx > 1
    ]
    if marker_rows:
        justification_rows = list(range(1, min(marker_rows)))
    else:
        # Fallback to the original template assumption (rows 1..6).
        justification_rows = list(range(1, min(len(table.rows), 7)))

    if not justification_rows:
        return

    lines = [line.strip() for line in justification.splitlines() if line.strip()]
    if not lines and justification.strip():
        lines = [justification.strip()]

    for slot, row_idx in enumerate(justification_rows):
        row = table.rows[row_idx]
        if not row.cells:
            continue
        text = lines[slot] if slot < len(lines) else ""
        _set_cell_text(row.cells[0], text)


def _fill_impact_rows(table: Any, impact_deliverables: str, impact_delivery_date: str) -> None:
    deliverables_value = _normalize_yes_no(impact_deliverables)
    delivery_value = _normalize_yes_no(impact_delivery_date)

    deliverables_idx = _find_row_index_containing(table, "IMPACT OF CHANGE TO DELIVERABLES")
    if deliverables_idx is not None:
        row = table.rows[deliverables_idx]
        for idx in range(3, len(row.cells)):
            _set_cell_text(row.cells[idx], deliverables_value)

    delivery_idx = _find_row_index_containing(table, "IMPACT ON DELIVERY DATE")
    if delivery_idx is not None:
        row = table.rows[delivery_idx]
        for idx in range(2, len(row.cells)):
            _set_cell_text(row.cells[idx], delivery_value)


def _fill_payment_terms_row(table: Any, payment_terms: str) -> None:
    payment_idx = _find_row_index_containing(table, "PAYMENT TERMS")
    if payment_idx is None:
        return
    row = table.rows[payment_idx]
    for idx in range(2, len(row.cells)):
        _set_cell_text(row.cells[idx], payment_terms)


def _fill_currency_cell(table: Any, currency: str) -> None:
    total_idx = _find_row_index_containing(table, "Total (Excluding Taxes)")
    if total_idx is None:
        return
    row = table.rows[total_idx]
    if len(row.cells) > 6:
        _set_cell_text(row.cells[6], currency)


def _fill_line_item_rows(table: Any, line_items: list[dict[str, Any]]) -> None:
    if len(table.rows) < 13:
        return

    header_idx = _find_row_index_containing(table, "Qty. Req")
    if header_idx is None:
        return
    template_idx = header_idx + 1
    total_idx = _find_row_index_containing(table, "Total (Excluding Taxes)")
    if total_idx is None or total_idx <= template_idx:
        return

    # Keep at least one editable row.
    clean_items = [
        {
            "qty": _to_text(item.get("qty")),
            "reqDescription": _to_text(item.get("reqDescription")),
            "unitCost": _to_text(item.get("unitCost")),
            "selectedItems": _to_text(item.get("selectedItems")),
        }
        for item in line_items
        if any(_to_text(item.get(key)) for key in ("qty", "reqDescription", "unitCost", "selectedItems"))
    ]
    required_rows = max(1, len(clean_items))
    existing_rows = total_idx - template_idx

    if required_rows > existing_rows:
        template_row = table.rows[template_idx]._tr
        insert_before = table.rows[total_idx]._tr
        for _ in range(required_rows - existing_rows):
            cloned = deepcopy(template_row)
            insert_before.addprevious(cloned)
        total_idx = _find_row_index_containing(table, "Total (Excluding Taxes)") or total_idx
        existing_rows = total_idx - template_idx

    # Fill visible line-item rows and clear unused ones.
    for offset in range(existing_rows):
        row = table.rows[template_idx + offset]
        payload = clean_items[offset] if offset < len(clean_items) else {
            "qty": "",
            "reqDescription": "",
            "unitCost": "",
            "selectedItems": "",
        }
        _set_cor_row_values(row, payload)

    # Update total selected items value.
    total_selected = sum(_to_float(item.get("selectedItems")) for item in clean_items)
    if total_idx < len(table.rows) and len(table.rows[total_idx].cells) > 5:
        _set_cell_text(table.rows[total_idx].cells[5], _format_amount(total_selected))


def _set_cor_row_values(row: Any, payload: dict[str, str]) -> None:
    if len(row.cells) < 6:
        return
    # Template columns:
    # col1 qty | col2+3 description (merged) | col4 unit cost | col5 selected items
    _set_cell_text(row.cells[1], payload.get("qty", ""))
    _set_cell_text(row.cells[2], payload.get("reqDescription", ""))
    _set_cell_text(row.cells[4], payload.get("unitCost", ""))
    _set_cell_text(row.cells[5], payload.get("selectedItems", ""))


def _find_row_index_containing(table: Any, needle: str) -> int | None:
    target = needle.strip().lower()
    for idx, row in enumerate(table.rows):
        row_text = " ".join(cell.text for cell in row.cells).strip().lower()
        if target in row_text:
            return idx
    return None


def _fill_comments_rows(document: Document, comments: str) -> None:
    for table in document.tables:
        comments_idx = _find_row_index_containing(table, "COMMENTS:")
        if comments_idx is None:
            continue
        target_idx = comments_idx + 1 if comments_idx + 1 < len(table.rows) else comments_idx
        row = table.rows[target_idx]
        if not row.cells:
            continue
        _set_cell_text(row.cells[0], comments)
        for cell in row.cells[1:]:
            cell_text = cell.text or ""
            if "\u00abComments\u00bb" in cell_text or re.search(r"\{\{\s*Comments\s*\}\}", cell_text):
                _set_cell_text(cell, comments)


def _set_cell_text(cell: Any, value: str) -> None:
    text = _to_text(value)
    if cell.paragraphs:
        _set_paragraph_text_with_font_size(cell.paragraphs[0], text)
        for paragraph in cell.paragraphs[1:]:
            paragraph.text = ""
    else:
        cell.text = text
        if cell.paragraphs:
            _set_paragraph_text_with_font_size(cell.paragraphs[0], text)


def _set_paragraph_text_with_font_size(paragraph: Any, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        paragraph.runs[0].font.size = Pt(INSERTED_TEXT_FONT_SIZE_PT)
        for run in paragraph.runs[1:]:
            run.text = ""
        return
    run = paragraph.add_run(text)
    run.font.size = Pt(INSERTED_TEXT_FONT_SIZE_PT)


def _format_amount(value: float) -> str:
    if value <= 0:
        return "0"
    return f"{value:,.2f}"


def _cleanup_or_phrase(value: str) -> str:
    text = _to_text(value)
    if not text:
        return ""
    text = re.sub(r"^\s*or\b\s*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s*\bor\b\s*$", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s{2,}", " ", text).strip()
    return text


def _resolve_initiator_value(contact_person: str, capmatic_pm: str, source: str) -> str:
    normalized_source = _to_text(source).lower()
    if normalized_source == "capmatic_pm":
        return _to_text(capmatic_pm)
    return _to_text(contact_person)


def _replace_initiator_pair(value: str, initiator_value: str) -> str:
    text = value
    replacement = _to_text(initiator_value)
    brace_pattern = r"\{\{\s*Contact_Person\s*\}\}\s*or\s*\{\{\s*Capmatic_PM\s*\}\}"
    angle_pattern = r"\u00ab\s*Contact_Person\s*\u00bb\s*or\s*\u00ab\s*Capmatic_PM\s*\u00bb"
    text = re.sub(brace_pattern, replacement, text, flags=re.IGNORECASE)
    text = re.sub(angle_pattern, replacement, text, flags=re.IGNORECASE)
    return text


def _normalize_yes_no(value: str) -> str:
    text = _to_text(value).lower()
    if text in {"yes", "y", "true", "1"}:
        return "Yes"
    if text in {"no", "n", "false", "0"}:
        return "No"
    return ""


def _format_cor_no_label(cor_no: str) -> str:
    value = _to_text(cor_no)
    if not value:
        return "COR"
    if re.match(r"^cor\b", value, flags=re.IGNORECASE):
        normalized = re.sub(r"^cor\b\.?\s*", "", value, flags=re.IGNORECASE)
        return f"COR {normalized}".strip()
    return f"COR {value}"
