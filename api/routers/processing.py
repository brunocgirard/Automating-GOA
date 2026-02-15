"""Processing workflow endpoints."""

from __future__ import annotations

import json
import os
import re
import sqlite3
from datetime import datetime
from functools import lru_cache
from html import escape
from typing import Any

from fastapi import APIRouter, HTTPException, Query, status
from fastapi.responses import FileResponse
from docx import Document

from api.models.schemas import (
    ExtractionRequest,
    ExtractionResponse,
    FillFormRequest,
    FillFormResponse,
    GenerateDocumentRequest,
    GenerateDocumentResponse,
    GoaFormDetailResponse,
    GoaGenerateDocumentApiRequest,
    GoaGenerateDocumentApiResponse,
    GoaGenerateDocumentRequest,
    GoaFormListItemResponse,
    GoaOutputOptions,
    GoaFormSchemaResponse,
    GoaFormSaveRequest,
    GoaFormSaveResponse,
    IdentifyMachinesRequest,
    MachineGroupingRequest,
    MachineGroupingResponse,
    MachineProcessingDataResponse,
    ProcessingArtifactsResponse,
    PricedItemResponse,
)
from api.services.processing_service import (
    build_options_listing,
    generate_document,
    get_machine_id_from_data,
    load_machine_by_id,
    load_quote_artifacts,
    run_extraction,
    save_generated_template,
)
from src.utils.db import (
    DB_PATH,
    group_items_by_confirmed_machines,
    load_document_content,
    load_goa_modifications,
    load_machines_for_quote,
    load_machine_template_data,
    load_priced_items_for_quote,
    save_machines_data,
    save_bulk_goa_modifications,
    save_machine_template_data,
)
from src.utils import template_utils
from src.utils.doc_filler import fill_word_document_from_llm_data
from src.utils.form_generator import OUTPUT_HTML_PATH, display_label, generate_goa_form, load_rows
from src.utils.html_doc_filler import fill_and_generate_html, fill_html_template
from src.utils.pdf_utils import identify_machines_from_items

router = APIRouter(prefix="/api/processing", tags=["Processing"])
direct_router = APIRouter(prefix="/api", tags=["Processing"])

SORTSTAR_TEMPLATE_CANDIDATES = (
    os.path.join("templates", "GOA_Sortstar_Temp.docx"),
    os.path.join("templates", "goa_sortstar_temp.docx"),
)
DEFAULT_HIDE_EMPTY_SECTIONS = True
DEFAULT_HIDE_EMPTY_FIELDS = True


def _normalize_template_data(raw_data: dict[str, Any]) -> dict[str, str]:
    normalized: dict[str, str] = {}
    for key, value in raw_data.items():
        if value is None:
            normalized[key] = ""
        elif isinstance(value, bool):
            normalized[key] = "YES" if value else "NO"
        else:
            normalized[key] = str(value)
    return normalized


def _slugify_identifier(value: str) -> str:
    normalized = re.sub(r"[^a-z0-9]+", "_", (value or "").strip().lower()).strip("_")
    return normalized or "section"


def _ensure_goa_form_template_path() -> str:
    template_path = str(OUTPUT_HTML_PATH)
    if not os.path.exists(template_path) and not generate_goa_form():
        raise RuntimeError("Failed to generate GOA HTML form template from Excel.")
    return template_path


def _render_goa_form_html(
    filled_data: dict[str, Any],
    *,
    hide_empty_sections: bool = False,
    hide_empty_fields: bool = False,
    included_sections: list[str] | None = None,
    label_overrides: dict[str, str] | None = None,
    pure_output: bool = False,
) -> str:
    template_path = _ensure_goa_form_template_path()
    with open(template_path, "r", encoding="utf-8") as file_handle:
        template_html = file_handle.read()
    return fill_html_template(
        template_html,
        _normalize_template_data(filled_data),
        hide_empty_sections=hide_empty_sections,
        hide_empty_fields=hide_empty_fields,
        included_sections=included_sections,
        label_overrides=label_overrides,
        pure_output=pure_output,
    )


def _normalize_section_title(section_name: str) -> str:
    cleaned = re.sub(r"\s*\(section\)\s*", "", section_name or "", flags=re.IGNORECASE).strip()
    return cleaned or "General"


def _resolve_schema_field_type(raw_type: str, placeholder: str) -> str:
    normalized = (raw_type or "").strip().lower()
    if placeholder == "options_listing" or normalized == "textarea":
        return "textarea"
    if normalized == "checkbox":
        return "checkbox"
    if normalized in {"qty", "number", "numeric"}:
        return "number"
    return "text"


@lru_cache(maxsize=1)
def _load_goa_form_schema() -> dict[str, Any]:
    rows = load_rows()
    sections: list[dict[str, Any]] = []
    section_index: dict[str, dict[str, Any]] = {}
    section_slug_counts: dict[str, int] = {}
    active_group_keys: dict[str, tuple[str, str] | None] = {}
    field_count = 0

    for row in rows:
        key = str(row.get("placeholder", "")).strip()
        if not key:
            continue

        section_title = _normalize_section_title(str(row.get("section", "")))
        subsection = str(row.get("subsection", "")).strip()
        subsub = str(row.get("subsub", "")).strip()
        group_key = (subsection, subsub)

        section = section_index.get(section_title)
        if section is None:
            base_slug = _slugify_identifier(section_title)
            occurrence = section_slug_counts.get(base_slug, 0)
            section_id = base_slug if occurrence == 0 else f"{base_slug}_{occurrence + 1}"
            section_slug_counts[base_slug] = occurrence + 1
            section = {"id": section_id, "title": section_title, "field_count": 0, "groups": []}
            section_index[section_title] = section
            sections.append(section)
            active_group_keys[section_title] = None

        if active_group_keys.get(section_title) != group_key:
            title_parts = [part for part in group_key if part]
            group_title = " / ".join(title_parts) if title_parts else None
            section["groups"].append({"title": group_title, "fields": []})
            active_group_keys[section_title] = group_key

        field_label = display_label(str(row.get("field", "")).strip()) or key
        section["groups"][-1]["fields"].append(
            {
                "key": key,
                "label": field_label,
                "type": _resolve_schema_field_type(str(row.get("type", "")), key),
            }
        )
        section["field_count"] += 1
        field_count += 1

    return {
        "sections": sections,
        "field_count": field_count,
    }


def _sanitize_machine_name(machine_name: str) -> str:
    safe = re.sub(r'[\\/*?:"<>|]', "_", machine_name.replace(" ", "_"))
    return safe or "machine"


def _is_sortstar_machine(machine_name: str) -> bool:
    return bool(re.search(r"\b(sortstar|unscrambler|bottle unscrambler)\b", (machine_name or "").lower()))


def _is_sortstar_template_row(template_row: dict[str, Any]) -> bool:
    template_type = str(template_row.get("template_type", "")).lower()
    machine_name = str(template_row.get("machine_name", "")).lower()
    generated_file_path = str(template_row.get("generated_file_path", "")).lower()
    if "sortstar" in template_type:
        return True
    if generated_file_path.endswith(".docx"):
        return True
    return bool(re.search(r"\b(sortstar|unscrambler|bottle unscrambler)\b", machine_name))


def _resolve_sortstar_template_path() -> str:
    for candidate in SORTSTAR_TEMPLATE_CANDIDATES:
        if os.path.exists(candidate):
            return candidate
    return SORTSTAR_TEMPLATE_CANDIDATES[0]


def _resolve_output_path(existing_file_path: str | None, machine_id: int, machine_name: str) -> str:
    if existing_file_path and existing_file_path.lower().endswith((".html", ".docx", ".pdf")):
        return existing_file_path
    if _is_sortstar_machine(machine_name):
        return f"output_SORTSTAR_{_sanitize_machine_name(machine_name)}_GOA.docx"
    return f"output_{machine_id}_{_sanitize_machine_name(machine_name)}_GOA.html"


def _resolve_generated_file_abs_path(file_path: str | None) -> str | None:
    if not file_path:
        return None
    resolved = os.path.abspath(file_path)
    if os.path.exists(resolved):
        return resolved
    return None


def _resolve_output_options(options: GoaOutputOptions | None) -> GoaOutputOptions:
    if options is None:
        return GoaOutputOptions(
            hide_empty_sections=DEFAULT_HIDE_EMPTY_SECTIONS,
            hide_empty_fields=DEFAULT_HIDE_EMPTY_FIELDS,
            pure_output=False,
            label_overrides={},
            format="html",
        )

    return GoaOutputOptions(
        included_sections=[str(value) for value in (options.included_sections or []) if str(value).strip()]
        or None,
        hide_empty_sections=bool(options.hide_empty_sections),
        hide_empty_fields=bool(options.hide_empty_fields),
        pure_output=bool(options.pure_output),
        label_overrides={str(key): str(value) for key, value in (options.label_overrides or {}).items()},
        format="html",
    )


def _serialize_output_options(options: GoaOutputOptions | None) -> dict[str, Any] | None:
    if options is None:
        return None
    resolved = _resolve_output_options(options)
    return {
        "included_sections": resolved.included_sections,
        "hide_empty_sections": resolved.hide_empty_sections,
        "hide_empty_fields": resolved.hide_empty_fields,
        "pure_output": resolved.pure_output,
        "label_overrides": resolved.label_overrides,
        "format": resolved.format,
    }


def _load_saved_output_options(template_row: dict[str, Any]) -> GoaOutputOptions | None:
    raw_preferences = template_row.get("output_preferences")
    if not isinstance(raw_preferences, dict):
        return None
    try:
        parsed = GoaOutputOptions(**raw_preferences)
    except Exception:
        return None
    return _resolve_output_options(parsed)


def _render_sortstar_docx_preview_html(file_path: str | None) -> str:
    resolved_path = _resolve_generated_file_abs_path(file_path)
    if not resolved_path:
        return (
            "<div style='padding:12px;border:1px solid #e5e7eb;border-radius:8px;background:#f9fafb;'>"
            "No generated DOCX file found for preview."
            "</div>"
        )

    try:
        document = Document(resolved_path)
    except Exception as exc:  # pragma: no cover - defensive guard
        return (
            "<div style='padding:12px;border:1px solid #fecaca;border-radius:8px;background:#fef2f2;color:#991b1b;'>"
            f"Failed to load DOCX preview: {escape(str(exc))}"
            "</div>"
        )

    paragraph_html: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            paragraph_html.append(f"<p>{escape(text)}</p>")

    table_html: list[str] = []
    for table in document.tables:
        table_html.append("<table>")
        for row in table.rows:
            table_html.append("<tr>")
            for cell in row.cells:
                cell_text = cell.text.strip()
                table_html.append(f"<td>{escape(cell_text)}</td>")
            table_html.append("</tr>")
        table_html.append("</table>")

    body_parts = paragraph_html + table_html
    if not body_parts:
        body_parts.append("<p>No text content found in document.</p>")

    body_html = "\n".join(body_parts)
    return f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <style>
    body {{
      font-family: "Segoe UI", Arial, sans-serif;
      margin: 0;
      padding: 16px;
      color: #111827;
      background: #ffffff;
      line-height: 1.4;
    }}
    p {{
      margin: 0 0 10px 0;
      white-space: pre-wrap;
    }}
    table {{
      border-collapse: collapse;
      width: 100%;
      margin: 14px 0;
      table-layout: fixed;
      font-size: 13px;
    }}
    td {{
      border: 1px solid #d1d5db;
      padding: 8px;
      vertical-align: top;
      white-space: pre-wrap;
      word-wrap: break-word;
    }}
  </style>
</head>
<body>
  {body_html}
</body>
</html>"""


def _build_field_labels(template_row: dict[str, Any], template_data: dict[str, Any]) -> dict[str, str]:
    if not template_data:
        return {}

    is_sortstar = _is_sortstar_template_row(template_row)
    mapping = (
        template_utils.SORTSTAR_EXPLICIT_MAPPINGS
        if is_sortstar
        else template_utils.DEFAULT_EXPLICIT_MAPPINGS
    )
    labels: dict[str, str] = {}
    for key in template_data.keys():
        mapped = mapping.get(key)
        if isinstance(mapped, str) and mapped.strip():
            labels[key] = mapped.strip()
    return labels


def _load_goa_template_row(machine_template_id: int) -> dict[str, Any] | None:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    try:
        cursor = conn.cursor()
        cursor.execute(
            """
            SELECT
                mt.id AS machine_template_id,
                mt.machine_id,
                mt.template_type,
                mt.template_data_json,
                mt.output_preferences_json,
                mt.generated_file_path,
                mt.processing_date,
                m.machine_name,
                m.client_quote_ref AS quote_ref
            FROM machine_templates mt
            JOIN machines m ON m.id = mt.machine_id
            WHERE mt.id = ?
            """,
            (machine_template_id,),
        )
        row = cursor.fetchone()
        if not row:
            return None
        payload = dict(row)
        template_json = payload.get("template_data_json") or "{}"
        try:
            payload["template_data"] = json.loads(template_json)
        except json.JSONDecodeError:
            payload["template_data"] = {}

        output_preferences_json = payload.get("output_preferences_json") or "{}"
        try:
            parsed_output_preferences = json.loads(output_preferences_json)
            payload["output_preferences"] = (
                parsed_output_preferences if isinstance(parsed_output_preferences, dict) else {}
            )
        except json.JSONDecodeError:
            payload["output_preferences"] = {}
        return payload
    finally:
        conn.close()


def _list_goa_templates(quote_ref: str | None = None, machine_id: int | None = None) -> list[dict[str, Any]]:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    try:
        cursor = conn.cursor()
        query = """
            SELECT
                mt.id AS machine_template_id,
                mt.machine_id,
                mt.template_type,
                mt.generated_file_path,
                mt.processing_date,
                m.machine_name,
                m.client_quote_ref AS quote_ref
            FROM machine_templates mt
            JOIN machines m ON m.id = mt.machine_id
            WHERE 1=1
        """
        params: list[Any] = []

        if quote_ref:
            query += " AND m.client_quote_ref = ?"
            params.append(quote_ref)
        if machine_id:
            query += " AND mt.machine_id = ?"
            params.append(machine_id)

        query += " ORDER BY mt.processing_date DESC"
        cursor.execute(query, params)
        return [dict(row) for row in cursor.fetchall()]
    finally:
        conn.close()


def _generate_goa_output(
    template_row: dict[str, Any],
    filled_data: dict[str, str],
    *,
    output_options: GoaOutputOptions | None = None,
) -> tuple[str, bool]:
    resolved_options = _resolve_output_options(output_options)
    is_sortstar_template = _is_sortstar_template_row(template_row)
    output_path = _resolve_output_path(
        template_row.get("generated_file_path"),
        int(template_row["machine_id"]),
        str(template_row.get("machine_name", "machine")),
    )

    if is_sortstar_template:
        docx_path = output_path
        if not docx_path.lower().endswith(".docx"):
            docx_path = os.path.splitext(docx_path)[0] + ".docx"

        fill_word_document_from_llm_data(_resolve_sortstar_template_path(), filled_data, docx_path)

        return docx_path, is_sortstar_template

    html_path = output_path
    if not html_path.lower().endswith(".html"):
        html_path = os.path.splitext(html_path)[0] + ".html"

    fill_and_generate_html(
        _ensure_goa_form_template_path(),
        filled_data,
        html_path,
        hide_empty_sections=resolved_options.hide_empty_sections,
        hide_empty_fields=resolved_options.hide_empty_fields,
        included_sections=resolved_options.included_sections,
        label_overrides=resolved_options.label_overrides,
        pure_output=resolved_options.pure_output,
    )

    return html_path, is_sortstar_template


@router.post("/identify", response_model=MachineGroupingResponse)
def identify_machines(payload: IdentifyMachinesRequest) -> dict:
    grouped = identify_machines_from_items(payload.items)
    return {"machines": grouped.get("machines", []), "common_items": grouped.get("common_items", [])}


@router.post("/group", response_model=MachineGroupingResponse)
def group_items(payload: MachineGroupingRequest) -> dict:
    grouped = group_items_by_confirmed_machines(
        payload.all_items,
        payload.main_machine_indices,
        payload.common_option_indices,
    )
    if payload.quote_ref:
        saved = save_machines_data(payload.quote_ref, grouped)
        if not saved:
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail="Failed to save grouped machine data.",
            )

        persisted_machines: list[dict[str, Any]] = []
        for machine_row in load_machines_for_quote(payload.quote_ref):
            machine_payload = machine_row.get("machine_data")
            if not isinstance(machine_payload, dict):
                continue
            persisted_machine = dict(machine_payload)
            persisted_machine["id"] = machine_row.get("id")
            persisted_machine.setdefault("machine_name", machine_row.get("machine_name", "Machine"))
            persisted_machines.append(persisted_machine)

        return {
            "machines": persisted_machines,
            "common_items": grouped.get("common_items", []),
        }
    return {"machines": grouped.get("machines", []), "common_items": grouped.get("common_items", [])}


@router.get("/items/{quote_ref}", response_model=list[PricedItemResponse])
def get_quote_items(quote_ref: str) -> list[dict]:
    return load_priced_items_for_quote(quote_ref)


@router.get("/artifacts/{quote_ref}", response_model=ProcessingArtifactsResponse)
def get_quote_artifacts(quote_ref: str) -> dict:
    data = load_quote_artifacts(quote_ref)
    return {
        "quote_ref": data["quote_ref"],
        "full_pdf_text": data.get("full_pdf_text", ""),
        "pdf_filename": data.get("pdf_filename"),
        "items": data.get("items", []),
        "machines": data.get("machines", []),
    }


@router.get("/machine-data/{machine_id}", response_model=MachineProcessingDataResponse)
def get_machine_data(machine_id: int) -> dict[str, Any]:
    machine_row = load_machine_by_id(machine_id)
    if not machine_row:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Machine not found.")

    machine_data = machine_row.get("machine_data")
    if not isinstance(machine_data, dict):
        machine_data = {}

    normalized_machine_data = dict(machine_data)
    normalized_machine_data["id"] = machine_id
    if not normalized_machine_data.get("machine_name"):
        normalized_machine_data["machine_name"] = machine_row.get("machine_name", "Machine")

    main_item = normalized_machine_data.get("main_item")
    if not isinstance(main_item, dict):
        main_item = {}
    options = normalized_machine_data.get("add_ons")
    if not isinstance(options, list):
        options = []
    common_items = normalized_machine_data.get("common_items")
    if not isinstance(common_items, list):
        common_items = []

    quote_ref = machine_row.get("client_quote_ref", "")
    document = load_document_content(quote_ref) or {}
    full_pdf_text = document.get("full_pdf_text", "") if isinstance(document, dict) else ""

    return {
        "machine_id": machine_id,
        "quote_ref": quote_ref,
        "machine_data": normalized_machine_data,
        "main_item": main_item,
        "options": options,
        "common_items": common_items,
        "full_pdf_text": full_pdf_text,
    }


@router.post("/extract", response_model=ExtractionResponse)
def extract_machine_fields(payload: ExtractionRequest) -> dict:
    try:
        result = run_extraction(
            machine_data=payload.machine_data,
            common_items=payload.common_items,
            template_contexts=payload.template_contexts,
            full_pdf_text=payload.full_pdf_text,
        )
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Machine extraction failed: {exc}",
        ) from exc

    return {
        "filled_data": result["filled_data"],
        "confidence_scores": result["confidence_scores"],
        "suggestions": result["suggestions"],
    }


@router.post("/generate", response_model=GenerateDocumentResponse)
def generate_machine_document(payload: GenerateDocumentRequest) -> dict:
    try:
        common_items = payload.common_items or payload.machine_data.get("common_items", []) or []
        file_path, _ = generate_document(
            machine_data=payload.machine_data,
            filled_data=payload.filled_data,
            common_items=common_items,
        )

        machine_id = payload.machine_id or get_machine_id_from_data(payload.machine_data)
        machine_template_id: int | None = None
        if machine_id:
            template_payload = dict(payload.filled_data)
            template_payload["options_listing"] = build_options_listing(
                payload.machine_data,
                common_items,
                template_payload,
            )
            save_generated_template(machine_id, template_payload, file_path, template_type="GOA")
            template_row = load_machine_template_data(machine_id, "GOA")
            if template_row:
                machine_template_id = int(template_row["id"])

        return {
            "file_path": file_path,
            "machine_id": machine_id,
            "machine_template_id": machine_template_id,
        }
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Document generation failed: {exc}",
        ) from exc


@router.post("/fill-form", response_model=FillFormResponse)
def fill_goa_form(payload: FillFormRequest) -> dict[str, str]:
    try:
        output_options = (
            _resolve_output_options(payload.output_options)
            if payload.output_options is not None
            else GoaOutputOptions(
                hide_empty_sections=False,
                hide_empty_fields=False,
                pure_output=False,
                label_overrides={},
                format="html",
            )
        )
        html = _render_goa_form_html(
            payload.filled_data,
            hide_empty_sections=output_options.hide_empty_sections,
            hide_empty_fields=output_options.hide_empty_fields,
            included_sections=output_options.included_sections,
            label_overrides=output_options.label_overrides,
            pure_output=output_options.pure_output,
        )
        return {"html": html}
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to render GOA form: {exc}",
        ) from exc


@router.get("/goa-schema", response_model=GoaFormSchemaResponse)
def get_goa_form_schema() -> dict[str, Any]:
    try:
        return _load_goa_form_schema()
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to load GOA form schema: {exc}",
        ) from exc


@router.get("/goa-form/{machine_template_id}", response_model=GoaFormDetailResponse)
def get_saved_goa_form(machine_template_id: int) -> dict[str, Any]:
    template_row = _load_goa_template_row(machine_template_id)
    if not template_row:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="GOA form not found.")

    template_data = dict(template_row.get("template_data", {}))
    modifications = load_goa_modifications(machine_template_id)
    for change in modifications:
        key = change.get("field_key")
        if key:
            template_data[key] = change.get("modified_value", "")

    is_sortstar_template = _is_sortstar_template_row(template_row)
    saved_output_options = _load_saved_output_options(template_row)
    field_labels = _build_field_labels(template_row, template_data)
    html = (
        _render_sortstar_docx_preview_html(template_row.get("generated_file_path"))
        if is_sortstar_template
        else _render_goa_form_html(
            template_data,
            hide_empty_sections=saved_output_options.hide_empty_sections if saved_output_options else False,
            hide_empty_fields=saved_output_options.hide_empty_fields if saved_output_options else False,
            included_sections=saved_output_options.included_sections if saved_output_options else None,
            label_overrides=saved_output_options.label_overrides if saved_output_options else None,
            pure_output=saved_output_options.pure_output if saved_output_options else False,
        )
    )

    return {
        "machine_template_id": machine_template_id,
        "machine_id": template_row["machine_id"],
        "quote_ref": template_row["quote_ref"],
        "machine_name": template_row["machine_name"],
        "template_type": template_row["template_type"],
        "generated_file_path": template_row.get("generated_file_path"),
        "processing_date": template_row.get("processing_date"),
        "template_data": template_data,
        "output_options": _serialize_output_options(saved_output_options),
        "field_labels": field_labels,
        "modifications": modifications,
        "html": html,
    }


@router.put("/goa-form/{machine_template_id}", response_model=GoaFormSaveResponse)
def save_goa_form(machine_template_id: int, payload: GoaFormSaveRequest) -> dict[str, Any]:
    template_row = _load_goa_template_row(machine_template_id)
    if not template_row:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="GOA form not found.")

    saved_output_options = _load_saved_output_options(template_row)
    requested_output_options = (
        _resolve_output_options(payload.output_options)
        if payload.output_options is not None
        else None
    )

    existing_data = _normalize_template_data(template_row.get("template_data", {}))
    incoming_data = _normalize_template_data(payload.filled_data)
    filled_data = dict(existing_data)
    filled_data.update(incoming_data)
    persisted_file_path = str(template_row.get("generated_file_path") or "")

    if not save_machine_template_data(
        machine_id=template_row["machine_id"],
        template_type=template_row["template_type"],
        template_data=filled_data,
        generated_file_path=persisted_file_path,
        output_preferences=_serialize_output_options(requested_output_options),
    ):
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Failed to save GOA form template data.",
        )

    change_map: dict[str, dict[str, str]] = {}
    if payload.modifications:
        for modification in payload.modifications:
            original_value = modification.original_value
            if original_value is None:
                original_value = existing_data.get(modification.field_key, "")
            change_map[modification.field_key] = {
                "original_value": str(original_value),
                "new_value": str(modification.modified_value),
            }
    else:
        keys = set(existing_data.keys()) | set(filled_data.keys())
        for key in keys:
            old_value = existing_data.get(key, "")
            new_value = filled_data.get(key, "")
            if old_value != new_value:
                change_map[key] = {"original_value": old_value, "new_value": new_value}

    if change_map:
        if not save_bulk_goa_modifications(
            machine_template_id=machine_template_id,
            changes=change_map,
            modification_reason="GOA form edit",
            modified_by="Next.js UI",
            regenerate_document=False,
        ):
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail="Failed to save GOA form modifications.",
            )

    if requested_output_options is not None:
        response_output_options = requested_output_options
    elif saved_output_options is not None:
        response_output_options = saved_output_options
    elif payload.generate_output:
        response_output_options = _resolve_output_options(None)
    else:
        response_output_options = GoaOutputOptions(
            hide_empty_sections=False,
            hide_empty_fields=False,
            pure_output=False,
            label_overrides={},
            format="html",
        )
    output_path = persisted_file_path
    if payload.generate_output:
        try:
            output_path, _ = _generate_goa_output(
                template_row,
                filled_data,
                output_options=response_output_options,
            )
        except RuntimeError as exc:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to regenerate GOA document: {exc}",
            ) from exc
        except Exception as exc:
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"Failed to regenerate GOA document: {exc}",
            ) from exc

        if not save_machine_template_data(
            machine_id=template_row["machine_id"],
            template_type=template_row["template_type"],
            template_data=filled_data,
            generated_file_path=output_path,
            output_preferences=_serialize_output_options(requested_output_options),
        ):
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail="Failed to persist generated GOA document path.",
            )

    refreshed_row = _load_goa_template_row(machine_template_id)
    if not refreshed_row:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="GOA form saved but could not be reloaded.",
        )

    refreshed_data = dict(refreshed_row.get("template_data", {}))
    for change in load_goa_modifications(machine_template_id):
        key = change.get("field_key")
        if key:
            refreshed_data[key] = change.get("modified_value", "")

    is_sortstar_template = _is_sortstar_template_row(refreshed_row)
    html = (
        _render_sortstar_docx_preview_html(refreshed_row.get("generated_file_path"))
        if is_sortstar_template
        else _render_goa_form_html(
            refreshed_data,
            hide_empty_sections=response_output_options.hide_empty_sections,
            hide_empty_fields=response_output_options.hide_empty_fields,
            included_sections=response_output_options.included_sections,
            label_overrides=response_output_options.label_overrides,
            pure_output=response_output_options.pure_output,
        )
    )
    return {
        "machine_template_id": machine_template_id,
        "saved_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "file_path": output_path,
        "html": html,
    }


@router.post("/goa-form/{machine_template_id}/generate-document", response_model=GoaFormSaveResponse)
def generate_goa_form_document(
    machine_template_id: int,
    payload: GoaGenerateDocumentRequest | None = None,
) -> dict[str, Any]:
    template_row = _load_goa_template_row(machine_template_id)
    if not template_row:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="GOA form not found.")

    filled_data = _normalize_template_data(template_row.get("template_data", {}))
    for change in load_goa_modifications(machine_template_id):
        key = change.get("field_key")
        if key:
            filled_data[key] = str(change.get("modified_value", ""))

    saved_output_options = _load_saved_output_options(template_row)
    requested_output_options = (
        _resolve_output_options(payload.output_options)
        if payload and payload.output_options is not None
        else None
    )
    output_options = (
        requested_output_options
        or saved_output_options
        or _resolve_output_options(None)
    )

    try:
        output_path, is_sortstar_template = _generate_goa_output(
            template_row,
            filled_data,
            output_options=output_options,
        )
    except RuntimeError as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Failed to generate GOA document: {exc}",
        ) from exc
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to generate GOA document: {exc}",
        ) from exc

    if not save_machine_template_data(
        machine_id=template_row["machine_id"],
        template_type=template_row["template_type"],
        template_data=filled_data,
        generated_file_path=output_path,
        output_preferences=_serialize_output_options(requested_output_options),
    ):
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Failed to persist generated GOA document path.",
        )

    refreshed_row = _load_goa_template_row(machine_template_id)
    if not refreshed_row:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="GOA document generated but form could not be reloaded.",
        )

    refreshed_data = dict(refreshed_row.get("template_data", {}))
    for change in load_goa_modifications(machine_template_id):
        key = change.get("field_key")
        if key:
            refreshed_data[key] = change.get("modified_value", "")

    html = (
        _render_sortstar_docx_preview_html(refreshed_row.get("generated_file_path"))
        if is_sortstar_template
        else _render_goa_form_html(
            refreshed_data,
            hide_empty_sections=output_options.hide_empty_sections,
            hide_empty_fields=output_options.hide_empty_fields,
            included_sections=output_options.included_sections,
            label_overrides=output_options.label_overrides,
            pure_output=output_options.pure_output,
        )
    )
    return {
        "machine_template_id": machine_template_id,
        "saved_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "file_path": output_path,
        "html": html,
    }


@direct_router.post("/generate-document", response_model=GoaGenerateDocumentApiResponse)
def generate_document_with_options(payload: GoaGenerateDocumentApiRequest) -> dict[str, Any]:
    template_row = _load_goa_template_row(payload.machine_template_id)
    if not template_row:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="GOA form not found.")

    existing_data = _normalize_template_data(template_row.get("template_data", {}))
    incoming_data = _normalize_template_data(payload.filled_data)
    filled_data = dict(existing_data)
    filled_data.update(incoming_data)
    saved_output_options = _load_saved_output_options(template_row)
    requested_output_options = (
        _resolve_output_options(payload.options)
        if payload.options is not None
        else None
    )
    output_options = (
        requested_output_options
        or saved_output_options
        or _resolve_output_options(None)
    )

    try:
        output_path, _ = _generate_goa_output(
            template_row,
            filled_data,
            output_options=output_options,
        )
    except RuntimeError as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=str(exc),
        ) from exc
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to generate document: {exc}",
        ) from exc

    if not save_machine_template_data(
        machine_id=template_row["machine_id"],
        template_type=template_row["template_type"],
        template_data=filled_data,
        generated_file_path=output_path,
        output_preferences=_serialize_output_options(requested_output_options),
    ):
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Failed to persist generated document path.",
        )

    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    output_format = "docx" if output_path.lower().endswith(".docx") else "html"
    return {
        "machine_template_id": payload.machine_template_id,
        "generated_at": generated_at,
        "format": output_format,
        "file_path": output_path,
        "download_url": f"/api/processing/goa-form/{payload.machine_template_id}/file",
    }


@router.get("/goa-form/{machine_template_id}/file")
def download_goa_form_file(machine_template_id: int) -> FileResponse:
    template_row = _load_goa_template_row(machine_template_id)
    if not template_row:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="GOA form not found.")

    resolved_path = _resolve_generated_file_abs_path(template_row.get("generated_file_path"))
    if not resolved_path:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Generated GOA file not found.",
        )

    lower_path = resolved_path.lower()
    if lower_path.endswith(".docx"):
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif lower_path.endswith(".pdf"):
        media_type = "application/pdf"
    else:
        media_type = "text/html"
    return FileResponse(
        path=resolved_path,
        media_type=media_type,
        filename=os.path.basename(resolved_path),
    )


@router.get("/goa-forms", response_model=list[GoaFormListItemResponse])
def list_goa_forms(
    quote_ref: str | None = Query(default=None),
    machine_id: int | None = Query(default=None),
) -> list[dict[str, Any]]:
    rows = _list_goa_templates(quote_ref=quote_ref, machine_id=machine_id)
    return [
        {
            "machine_template_id": row["machine_template_id"],
            "machine_id": row["machine_id"],
            "quote_ref": row["quote_ref"],
            "machine_name": row["machine_name"],
            "template_type": row["template_type"],
            "generated_file_path": row.get("generated_file_path"),
            "processing_date": row.get("processing_date"),
        }
        for row in rows
    ]
