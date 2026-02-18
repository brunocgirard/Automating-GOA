"""COR workflow helpers for prefill and DOCX generation."""

from __future__ import annotations

import re
from copy import deepcopy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document

OUTPUT_ROOT = Path("data/generated/cor")
TEMPLATE_PATH = Path("Mail_merge/Customer_Oxxxx_COR.docx")
TOKEN_PATTERN = re.compile(r"\u00ab[^\u00bb]+\u00bb")


@dataclass
class GeneratedCorArtifact:
    path: Path
    filename: str
    media_type: str


def build_cor_prefill_data(
    quote: dict[str, Any],
) -> dict[str, Any]:
    """Build a prefilled COR state from quote info only.

    COR line rows must always be user-entered because this document
    is used to add/remove items, not mirror quote priced items.
    """
    line_items = [
        {
            "id": "cor-line-1",
            "qty": "",
            "reqDescription": "",
            "unitCost": "",
            "selectedItems": "",
        }
    ]

    return {
        "quoteId": int(quote.get("id") or 0),
        "quoteRef": _to_text(quote.get("quote_ref")),
        "client": {
            "company": _to_text(quote.get("company")) or _to_text(quote.get("customer_name")),
            "customerPO": _to_text(quote.get("customer_po")),
            "orderDate": _to_text(quote.get("order_date")),
            "ax": _to_text(quote.get("ax")),
            "ox": _to_text(quote.get("ox")),
            "machine": _to_text(quote.get("machine_model")),
        },
        "corNo": "1",
        "justificationForChange": "",
        "lineItems": line_items,
    }


def generate_cor_document(cor_data: dict[str, Any], quote_ref: str) -> GeneratedCorArtifact:
    """Generate COR DOCX document from COR state."""
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"Template not found: {TEMPLATE_PATH}")

    safe_quote_ref = _slugify(quote_ref) or "quote"
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_dir = OUTPUT_ROOT / safe_quote_ref / timestamp
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"{safe_quote_ref}_cor.docx"

    document = Document(str(TEMPLATE_PATH))
    _replace_core_tokens(document, cor_data)
    _fill_cor_template_tables(document, cor_data)
    document.save(str(output_path))

    return GeneratedCorArtifact(
        path=output_path,
        filename=output_path.name,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


def _replace_core_tokens(document: Document, cor_data: dict[str, Any]) -> None:
    client = cor_data.get("client") if isinstance(cor_data.get("client"), dict) else {}
    replacements = {
        "Company": _to_text(client.get("company")),
        "Customer_PO": _to_text(client.get("customerPO")),
        "Ax": _to_text(client.get("ax")),
        "Ox": _to_text(client.get("ox")),
        "Machine": _to_text(client.get("machine")),
    }

    for paragraph in document.paragraphs:
        _replace_tokens_in_paragraph(paragraph, replacements)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_tokens_in_paragraph(paragraph, replacements)


def _fill_cor_template_tables(document: Document, cor_data: dict[str, Any]) -> None:
    tables = document.tables
    if len(tables) < 3:
        return

    client = cor_data.get("client") if isinstance(cor_data.get("client"), dict) else {}
    cor_no = _to_text(cor_data.get("corNo"))
    order_date = _to_text(client.get("orderDate"))
    justification = _to_text(cor_data.get("justificationForChange"))
    line_items = [entry for entry in cor_data.get("lineItems", []) if isinstance(entry, dict)]

    # Table 0 has COR no cell.
    table0 = tables[0]
    if len(table0.rows) > 4 and len(table0.rows[4].cells) > 1:
        _set_cell_text(table0.rows[4].cells[1], cor_no)

    # Table 1 has order date.
    table1 = tables[1]
    if len(table1.rows) > 1 and len(table1.rows[1].cells) > 1:
        _set_cell_text(table1.rows[1].cells[1], order_date)

    table2 = tables[2]
    _fill_justification_rows(table2, justification)
    _fill_line_item_rows(table2, line_items)


def _fill_justification_rows(table: Any, justification: str) -> None:
    if len(table.rows) < 8:
        return

    # Rows 1..6 are the editable justification area in this template.
    justification_rows = list(range(1, 7))
    lines = [line.strip() for line in justification.splitlines() if line.strip()]
    if not lines and justification.strip():
        lines = [justification.strip()]

    for slot, row_idx in enumerate(justification_rows):
        row = table.rows[row_idx]
        if not row.cells:
            continue
        text = lines[slot] if slot < len(lines) else ""
        _set_cell_text(row.cells[0], text)


def _fill_line_item_rows(table: Any, line_items: list[dict[str, Any]]) -> None:
    if len(table.rows) < 13:
        return

    header_idx = _find_row_index_containing(table, "Qty. Req")
    if header_idx is None:
        return
    template_idx = header_idx + 1
    total_idx = _find_row_index_containing(table, "Total (Excluding Taxes)")
    if total_idx is None or total_idx <= template_idx:
        return

    # Keep at least one editable row.
    clean_items = [
        {
            "qty": _to_text(item.get("qty")),
            "reqDescription": _to_text(item.get("reqDescription")),
            "unitCost": _to_text(item.get("unitCost")),
            "selectedItems": _to_text(item.get("selectedItems")),
        }
        for item in line_items
        if any(_to_text(item.get(key)) for key in ("qty", "reqDescription", "unitCost", "selectedItems"))
    ]
    required_rows = max(1, len(clean_items))
    existing_rows = total_idx - template_idx

    if required_rows > existing_rows:
        template_row = table.rows[template_idx]._tr
        insert_before = table.rows[total_idx]._tr
        for _ in range(required_rows - existing_rows):
            cloned = deepcopy(template_row)
            insert_before.addprevious(cloned)
        total_idx = _find_row_index_containing(table, "Total (Excluding Taxes)") or total_idx
        existing_rows = total_idx - template_idx

    # Fill visible line-item rows and clear unused ones.
    for offset in range(existing_rows):
        row = table.rows[template_idx + offset]
        payload = clean_items[offset] if offset < len(clean_items) else {
            "qty": "",
            "reqDescription": "",
            "unitCost": "",
            "selectedItems": "",
        }
        _set_cor_row_values(row, payload)

    # Update total selected items value.
    total_selected = sum(_to_float(item.get("selectedItems")) for item in clean_items)
    if total_idx < len(table.rows) and len(table.rows[total_idx].cells) > 5:
        _set_cell_text(table.rows[total_idx].cells[5], _format_amount(total_selected))


def _set_cor_row_values(row: Any, payload: dict[str, str]) -> None:
    if len(row.cells) < 6:
        return
    # Template columns:
    # col1 qty | col2+3 description (merged) | col4 unit cost | col5 selected items
    _set_cell_text(row.cells[1], payload.get("qty", ""))
    _set_cell_text(row.cells[2], payload.get("reqDescription", ""))
    _set_cell_text(row.cells[4], payload.get("unitCost", ""))
    _set_cell_text(row.cells[5], payload.get("selectedItems", ""))


def _find_row_index_containing(table: Any, needle: str) -> int | None:
    target = needle.strip().lower()
    for idx, row in enumerate(table.rows):
        row_text = " ".join(cell.text for cell in row.cells).strip().lower()
        if target in row_text:
            return idx
    return None


def _replace_tokens_in_paragraph(paragraph: Any, replacements: dict[str, str]) -> None:
    raw_text = "".join(run.text for run in paragraph.runs)
    if not raw_text:
        raw_text = paragraph.text or ""
    if not raw_text:
        return

    updated = raw_text
    for token, value in replacements.items():
        updated = updated.replace(f"\u00ab{token}\u00bb", value)
    updated = TOKEN_PATTERN.sub("", updated)

    if updated == raw_text:
        return
    if paragraph.runs:
        paragraph.runs[0].text = updated
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = updated


def _set_cell_text(cell: Any, value: str) -> None:
    text = _to_text(value)
    if cell.paragraphs:
        cell.paragraphs[0].text = text
        for paragraph in cell.paragraphs[1:]:
            paragraph.text = ""
    else:
        cell.text = text


def _format_amount(value: float) -> str:
    if value <= 0:
        return "0"
    return f"{value:,.2f}"


def _to_text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _to_float(value: Any) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    cleaned = re.sub(r"[^0-9.\-]", "", str(value))
    if not cleaned:
        return 0.0
    try:
        return float(cleaned)
    except ValueError:
        return 0.0


def _slugify(value: str) -> str:
    normalized = re.sub(r"[^a-zA-Z0-9]+", "-", value.strip()).strip("-")
    return normalized.lower()
