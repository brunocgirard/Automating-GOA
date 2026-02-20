"""Shipping workflow helpers for prefill, HS lookup, and shipping document generation."""

from __future__ import annotations

from html import escape
import json
import re
import zipfile
from copy import deepcopy
from dataclasses import dataclass
from datetime import datetime
from functools import lru_cache
from pathlib import Path
from typing import Any

from docx import Document

HS_CODE_FILE = Path("src/data/hs_codes.json")
OUTPUT_ROOT = Path("data/generated/shipping")

PACKING_SLIP_TEMPLATE_CANDIDATES = (
    Path("Mail_merge/Paking Slip1.docx"),
    Path("Mail_merge/Paking Slip.docx"),
)

COMMERCIAL_INVOICE_TEMPLATE_CANDIDATES = (
    Path("Mail_merge/Commercial Invoice1.docx"),
    Path("Mail_merge/Commercial Invoice.docx"),
)

CERTIFICATE_ORIGIN_TEMPLATE_CANDIDATES = (
    Path("Mail_merge/CERTIFICATION OF ORIGIN_NAFTA1.docx"),
    Path("Mail_merge/CERTIFICATION OF ORIGIN_NAFTA.docx"),
)

TEMPLATE_PATHS = {
    "packing_slip": Path("Mail_merge/Paking Slip.docx"),
    "commercial_invoice": Path("Mail_merge/Commercial Invoice.docx"),
    "certificate_origin": Path("Mail_merge/CERTIFICATION OF ORIGIN_NAFTA.docx"),
}

TOKEN_PATTERN = re.compile(r"\u00ab[^\u00bb]+\u00bb")
DOUBLE_BRACE_TOKEN_PATTERN = re.compile(r"\{\{\s*[^{}]+\s*\}\}")

DEFAULT_META: dict[str, str] = {
    "brokerInfo": "",
    "blanketFrom": "",
    "blanketTo": "",
    "originCriterion": "B",
    "certifierName": "",
    "certifierTitle": "Project Manager",
    "certifierDate": "",
    "certifierContact": "",
    "totalInvoiceAmount": "",
    "countryOfOrigin": "Canada",
}


@dataclass
class GeneratedShippingArtifact:
    path: Path
    filename: str
    media_type: str


@lru_cache(maxsize=1)
def load_hs_code_patterns() -> list[dict[str, str]]:
    """Load HS code pattern data from seed JSON."""
    if not HS_CODE_FILE.exists():
        return []

    try:
        raw = json.loads(HS_CODE_FILE.read_text(encoding="utf-8"))
    except Exception:
        return []

    rows: list[dict[str, str]] = []
    for item in raw if isinstance(raw, list) else []:
        if not isinstance(item, dict):
            continue
        pattern = str(item.get("machine_pattern") or "").strip()
        hs_code = str(item.get("hs_code") or "").strip()
        if not pattern or not hs_code:
            continue
        rows.append({"machine_pattern": pattern, "hs_code": hs_code})

    rows.sort(key=lambda entry: len(_normalize_machine_text(entry["machine_pattern"])), reverse=True)
    return rows


def match_machine_hs_code(machine_name: str, fallback: str = "") -> str:
    """Best-effort machine name to HS code lookup."""
    normalized_name = _normalize_machine_text(machine_name)
    if not normalized_name:
        return fallback

    for entry in load_hs_code_patterns():
        normalized_pattern = _normalize_machine_text(entry["machine_pattern"])
        if normalized_pattern and normalized_pattern in normalized_name:
            return entry["hs_code"]

    return fallback


def build_shipping_prefill_data(
    quote: dict[str, Any],
    machine_rows: list[dict[str, Any]],
    line_item_rows: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    """Build a prefilled shipping state from quote, machines, and priced line items."""
    sold_to = _split_address_lines(quote.get("sold_to_address"))
    ship_to = _split_address_lines(quote.get("ship_to_address"))

    machines: list[dict[str, Any]] = []
    for index, machine_row in enumerate(machine_rows, start=1):
        machine_payload = machine_row.get("machine_data") if isinstance(machine_row, dict) else {}
        if not isinstance(machine_payload, dict):
            machine_payload = {}

        machine_name = _to_text(machine_row.get("machine_name"))
        model = _resolve_machine_model(machine_name, machine_payload)
        hs_code = match_machine_hs_code(
            f"{machine_name} {model}".strip(),
            fallback=_to_text(quote.get("hs_code")),
        )

        machines.append(
            {
                "id": f"machine-{machine_row.get('id', index)}",
                "machineId": machine_row.get("id"),
                "machineName": machine_name or f"Machine {index}",
                "model": model or machine_name or f"Machine {index}",
                "hsCode": hs_code,
                "serialNumber": _to_text(quote.get("serial_number")),
                "unitPrice": _resolve_machine_price(machine_payload),
                "truckId": "truck-1",
                "crates": [
                    {
                        "id": f"crate-{index}-1",
                        "lengthIn": "",
                        "widthIn": "",
                        "heightIn": "",
                        "weightLbs": "",
                    }
                ],
            }
        )

    if not machines:
        fallback_machine = _to_text(quote.get("machine_model")) or "Machine"
        machines.append(
            {
                "id": "machine-1",
                "machineId": None,
                "machineName": fallback_machine,
                "model": fallback_machine,
                "hsCode": match_machine_hs_code(fallback_machine, fallback=_to_text(quote.get("hs_code"))),
                "serialNumber": _to_text(quote.get("serial_number")),
                "unitPrice": 0,
                "truckId": "truck-1",
                "crates": [
                    {
                        "id": "crate-1-1",
                        "lengthIn": "",
                        "widthIn": "",
                        "heightIn": "",
                        "weightLbs": "",
                    }
                ],
            }
        )

    total_invoice_amount = sum(_to_float(machine.get("unitPrice")) for machine in machines)
    meta = dict(DEFAULT_META)
    if total_invoice_amount > 0:
        meta["totalInvoiceAmount"] = f"{total_invoice_amount:.2f}"

    line_item_options = _build_line_item_options(line_item_rows or [])

    return {
        "quoteId": int(quote.get("id") or 0),
        "quoteRef": _to_text(quote.get("quote_ref")),
        "client": {
            "quoteRef": _to_text(quote.get("quote_ref")),
            "company": _to_text(quote.get("company")),
            "customerName": _to_text(quote.get("customer_name")) or _to_text(quote.get("company")),
            "soldToAddress1": sold_to[0],
            "soldToAddress2": sold_to[1],
            "soldToAddress3": sold_to[2],
            "shipToAddress1": ship_to[0],
            "shipToAddress2": ship_to[1],
            "shipToAddress3": ship_to[2],
            "telephone": _to_text(quote.get("telephone")),
            "customerPO": _to_text(quote.get("customer_po")),
            "orderDate": _to_text(quote.get("order_date")),
            "ax": _to_text(quote.get("ax")),
            "ox": _to_text(quote.get("ox")),
            "via": _to_text(quote.get("via")),
            "incoterm": _to_text(quote.get("incoterm")),
            "taxId": _to_text(quote.get("tax_id")),
            "customerNumber": _to_text(quote.get("customer_number")),
            "clientContact": _to_text(quote.get("customer_contact_person")),
        },
        "machines": machines,
        "lineItemOptions": line_item_options,
        "trucks": [{"id": "truck-1", "name": "Truck 1"}],
        "meta": meta,
    }


def generate_shipping_documents(
    shipping_data: dict[str, Any],
    quote_ref: str,
    document_type: str = "all",
    output_format: str = "docx",
) -> GeneratedShippingArtifact:
    """Generate shipping docs in DOCX/HTML, or a ZIP for multi-file output."""
    requested = document_type or "all"
    if requested not in TEMPLATE_PATHS and requested != "all":
        raise ValueError(f"Unsupported document_type '{requested}'.")
    fmt = (output_format or "docx").lower().strip()
    if fmt not in {"docx", "html"}:
        raise ValueError(f"Unsupported output_format '{output_format}'.")

    safe_quote_ref = _slugify(quote_ref) or "quote"
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_dir = OUTPUT_ROOT / safe_quote_ref / timestamp
    output_dir.mkdir(parents=True, exist_ok=True)

    docs_to_generate = list(TEMPLATE_PATHS.keys()) if requested == "all" else [requested]
    truck_payloads = _build_truck_payloads(shipping_data)
    multiple_trucks = len(truck_payloads) > 1
    generated: dict[str, Path] = {}

    for truck_payload in truck_payloads:
        truck_suffix = f"-truck-{truck_payload['index']}" if multiple_trucks else ""
        truck_state = truck_payload["shipping_data"]
        for key in docs_to_generate:
            ext = "docx" if fmt == "docx" else "html"
            output_name = {
                "packing_slip": f"packing-slip{truck_suffix}.{ext}",
                "commercial_invoice": f"commercial-invoice{truck_suffix}.{ext}",
                "certificate_origin": f"certificate-of-origin{truck_suffix}.{ext}",
            }[key]
            output_path = output_dir / output_name
            if fmt == "docx":
                _generate_single_docx(key, truck_state, output_path)
            else:
                _generate_single_html(key, truck_state, output_path)
            generated[output_name] = output_path

    if requested == "all" or len(generated) > 1:
        zip_name = (
            f"{safe_quote_ref}_shipping_documents.zip"
            if requested == "all"
            else f"{safe_quote_ref}_{requested}_by_truck.zip"
        )
        zip_path = output_dir / zip_name
        with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for path in generated.values():
                archive.write(path, path.name)
        return GeneratedShippingArtifact(
            path=zip_path,
            filename=zip_path.name,
            media_type="application/zip",
        )

    single_path = next(iter(generated.values()))
    return GeneratedShippingArtifact(
        path=single_path,
        filename=single_path.name,
        media_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            if fmt == "docx"
            else "text/html; charset=utf-8"
        ),
    )


def _generate_single_docx(document_type: str, shipping_data: dict[str, Any], output_path: Path) -> None:
    template_path = _resolve_template_path(document_type)
    if not template_path.exists():
        raise FileNotFoundError(f"Template not found: {template_path}")

    document = Document(str(template_path))
    machine_rows = _machine_rows(shipping_data)
    _fill_repeating_machine_rows(document, machine_rows)

    token_replacements = _build_global_token_replacements(shipping_data, machine_rows)
    literal_replacements = _build_literal_replacements(shipping_data)
    _replace_tokens_in_document(document, token_replacements, literal_replacements)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def _resolve_template_path(document_type: str) -> Path:
    if document_type == "packing_slip":
        for candidate in PACKING_SLIP_TEMPLATE_CANDIDATES:
            if candidate.exists():
                return candidate
    if document_type == "commercial_invoice":
        for candidate in COMMERCIAL_INVOICE_TEMPLATE_CANDIDATES:
            if candidate.exists():
                return candidate
    if document_type == "certificate_origin":
        for candidate in CERTIFICATE_ORIGIN_TEMPLATE_CANDIDATES:
            if candidate.exists():
                return candidate
    return TEMPLATE_PATHS[document_type]


def _generate_single_html(document_type: str, shipping_data: dict[str, Any], output_path: Path) -> None:
    html = _render_shipping_html(document_type, shipping_data)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(html, encoding="utf-8")


def _render_shipping_html(document_type: str, shipping_data: dict[str, Any]) -> str:
    if document_type == "packing_slip":
        return _render_packing_slip_html(shipping_data)
    if document_type == "commercial_invoice":
        return _render_commercial_invoice_html(shipping_data)
    if document_type == "certificate_origin":
        return _render_certificate_origin_html(shipping_data)
    raise ValueError(f"Unsupported document_type '{document_type}'.")


def _render_packing_slip_html(shipping_data: dict[str, Any]) -> str:
    client = shipping_data.get("client") if isinstance(shipping_data.get("client"), dict) else {}
    machines = [machine for machine in shipping_data.get("machines", []) if isinstance(machine, dict)]
    trucks = [truck for truck in shipping_data.get("trucks", []) if isinstance(truck, dict)]
    truck_name_by_id = {
        _to_text(truck.get("id")): _to_text(truck.get("name")) or "Truck"
        for truck in trucks
    }
    quote_ref = _to_text(shipping_data.get("quoteRef"))

    total_crates = 0
    total_weight = 0.0
    for machine in machines:
        crates = machine.get("crates")
        if not isinstance(crates, list):
            continue
        total_crates += len(crates)
        for crate in crates:
            if isinstance(crate, dict):
                total_weight += _to_float(crate.get("weightLbs"))

    rows_html = "".join(
        (
            "<tr>"
            f"<td>{escape(_line(truck_name_by_id.get(_to_text(machine.get('truckId'))) or 'Truck'))}</td>"
            f"<td>{escape(_line(_to_text(machine.get('model')) or _to_text(machine.get('machineName'))))}</td>"
            f"<td>{escape(_line(_to_text(machine.get('serialNumber'))))}</td>"
            f"<td>{escape(_line(_to_text(machine.get('hsCode'))))}</td>"
            f"<td style=\"text-align:right\">{len(machine.get('crates') or [])}</td>"
            "</tr>"
        )
        for machine in machines
    )
    if not rows_html:
        rows_html = "<tr><td colspan=\"5\" style=\"text-align:center;color:#666\">No machines configured.</td></tr>"

    body = f"""
    <div class="doc-header">
      <h1>Packing Slip</h1>
      <p>Quote: {escape(_line(quote_ref))}</p>
    </div>
    <div class="grid two">
      <div class="panel">
        <h3>Sold To</h3>
        <p>{escape(_line(_to_text(client.get("customerName")) or _to_text(client.get("company"))))}</p>
        <p>{escape(_line(_to_text(client.get("soldToAddress1"))))}</p>
        <p>{escape(_line(_to_text(client.get("soldToAddress2"))))}</p>
        <p>{_html_lines(_to_text(client.get("soldToAddress3")))}</p>
      </div>
      <div class="panel">
        <h3>Ship To</h3>
        <p>{escape(_line(_to_text(client.get("customerName")) or _to_text(client.get("company"))))}</p>
        <p>{escape(_line(_to_text(client.get("shipToAddress1"))))}</p>
        <p>{escape(_line(_to_text(client.get("shipToAddress2"))))}</p>
        <p>{_html_lines(_to_text(client.get("shipToAddress3")))}</p>
      </div>
    </div>
    <table>
      <tbody>
        <tr><th>Customer PO</th><td>{escape(_line(_to_text(client.get("customerPO"))))}</td><th>Order Date</th><td>{escape(_line(_to_text(client.get("orderDate"))))}</td></tr>
        <tr><th>Incoterm</th><td>{escape(_line(_to_text(client.get("incoterm"))))}</td><th>Totals</th><td>{total_crates} crates / {total_weight:.1f} lbs</td></tr>
      </tbody>
    </table>
    <table>
      <thead><tr><th>Truck</th><th>Machine</th><th>Serial</th><th>HS</th><th style="text-align:right">Crates</th></tr></thead>
      <tbody>{rows_html}</tbody>
    </table>
    """
    return _wrap_shipping_html("Packing Slip", body)


def _render_commercial_invoice_html(shipping_data: dict[str, Any]) -> str:
    client = shipping_data.get("client") if isinstance(shipping_data.get("client"), dict) else {}
    meta = shipping_data.get("meta") if isinstance(shipping_data.get("meta"), dict) else {}
    machines = [machine for machine in shipping_data.get("machines", []) if isinstance(machine, dict)]
    quote_ref = _to_text(shipping_data.get("quoteRef"))

    manual_total = _to_float(meta.get("totalInvoiceAmount"))
    computed_total = sum(_to_float(machine.get("unitPrice")) for machine in machines)
    total = manual_total if manual_total > 0 else computed_total

    rows_html = "".join(
        (
            "<tr>"
            "<td>1</td>"
            "<td>"
            f"<div>{escape(_line(_to_text(machine.get('model')) or _to_text(machine.get('machineName'))))}</div>"
            f"<div class=\"muted\">Serial: {escape(_line(_to_text(machine.get('serialNumber'))))}</div>"
            f"<div class=\"muted\">HS: {escape(_line(_to_text(machine.get('hsCode'))))}</div>"
            "</td>"
            f"<td style=\"text-align:right\">{_usd(_to_float(machine.get('unitPrice')))}</td>"
            f"<td style=\"text-align:right\">{_usd(_to_float(machine.get('unitPrice')))}</td>"
            "</tr>"
        )
        for machine in machines
    )
    if not rows_html:
        rows_html = "<tr><td colspan=\"4\" style=\"text-align:center;color:#666\">No machines configured.</td></tr>"

    body = f"""
    <div class="doc-header">
      <h1>Commercial Invoice</h1>
      <p>Quote: {escape(_line(quote_ref))} | Currency: USD</p>
    </div>
    <div class="grid two">
      <div class="panel">
        <h3>Sold To</h3>
        <p>{escape(_line(_to_text(client.get("customerName")) or _to_text(client.get("company"))))}</p>
        <p>{escape(_line(_to_text(client.get("soldToAddress1"))))}</p>
        <p>{escape(_line(_to_text(client.get("soldToAddress2"))))}</p>
        <p>{_html_lines(_to_text(client.get("soldToAddress3")))}</p>
        <p>Tax ID: {escape(_line(_to_text(client.get("taxId"))))}</p>
      </div>
      <div class="panel">
        <h3>Ship To</h3>
        <p>{escape(_line(_to_text(client.get("customerName")) or _to_text(client.get("company"))))}</p>
        <p>{escape(_line(_to_text(client.get("shipToAddress1"))))}</p>
        <p>{escape(_line(_to_text(client.get("shipToAddress2"))))}</p>
        <p>{_html_lines(_to_text(client.get("shipToAddress3")))}</p>
      </div>
    </div>
    <table>
      <tbody>
        <tr><th>Customer PO</th><td>{escape(_line(_to_text(client.get("customerPO"))))}</td><th>Order date</th><td>{escape(_line(_to_text(client.get("orderDate"))))}</td></tr>
        <tr><th>Order number</th><td>{escape(_line(_to_text(client.get("ox"))))}</td><th>Broker</th><td>{escape(_line(_to_text(meta.get("brokerInfo"))))}</td></tr>
        <tr><th>Incoterm</th><td>{escape(_line(_to_text(client.get("incoterm"))))}</td><th>Invoice total</th><td>{_usd(total)}</td></tr>
      </tbody>
    </table>
    <table>
      <thead><tr><th>Qty</th><th>Description</th><th style="text-align:right">Unit (USD)</th><th style="text-align:right">Total (USD)</th></tr></thead>
      <tbody>{rows_html}</tbody>
    </table>
    <div class="totals">
      <div><span>Amount before tax</span><strong>{_usd(total)}</strong></div>
      <div><span>Total amount (USD)</span><strong>{_usd(total)}</strong></div>
    </div>
    """
    return _wrap_shipping_html("Commercial Invoice", body)


def _render_certificate_origin_html(shipping_data: dict[str, Any]) -> str:
    client = shipping_data.get("client") if isinstance(shipping_data.get("client"), dict) else {}
    meta = shipping_data.get("meta") if isinstance(shipping_data.get("meta"), dict) else {}
    machines = [machine for machine in shipping_data.get("machines", []) if isinstance(machine, dict)]
    quote_ref = _to_text(shipping_data.get("quoteRef"))

    rows_html = "".join(
        (
            "<tr>"
            f"<td>{escape(_line(_to_text(machine.get('model')) or _to_text(machine.get('machineName'))))}</td>"
            f"<td>{escape(_line(_to_text(machine.get('hsCode'))))}</td>"
            f"<td>{escape(_line(_to_text(meta.get('originCriterion')) or 'B'))}</td>"
            f"<td>{escape(_line(_to_text(meta.get('countryOfOrigin')) or 'Canada'))}</td>"
            "</tr>"
        )
        for machine in machines
    )
    if not rows_html:
        rows_html = "<tr><td colspan=\"4\" style=\"text-align:center;color:#666\">No machines configured.</td></tr>"

    body = f"""
    <div class="doc-header">
      <h1>Certificate of Origin (USMCA / CUSMA)</h1>
      <p>Quote: {escape(_line(quote_ref))}</p>
    </div>
    <div class="grid two">
      <div class="panel">
        <h3>Exporter</h3>
        <p>CAPMATIC LTD</p>
        <p>12180 ALBERT-HUDON</p>
        <p>MONTREAL, QUEBEC, CANADA H1G 3K7</p>
        <p>Telephone: 514-322-0062</p>
      </div>
      <div class="panel">
        <h3>Importer</h3>
        <p>{escape(_line(_to_text(client.get("customerName")) or _to_text(client.get("company"))))}</p>
        <p>{escape(_line(_to_text(client.get("soldToAddress1"))))}</p>
        <p>{escape(_line(_to_text(client.get("soldToAddress2"))))}</p>
        <p>{_html_lines(_to_text(client.get("soldToAddress3")))}</p>
        <p>Tax ID: {escape(_line(_to_text(client.get("taxId"))))}</p>
      </div>
    </div>
    <table>
      <thead><tr><th>Description of goods</th><th>HS tariff</th><th>Origin criterion</th><th>Country of origin</th></tr></thead>
      <tbody>{rows_html}</tbody>
    </table>
    <div class="grid two">
      <div class="panel"><h3>Blanket period</h3><p>From: {escape(_line(_to_text(meta.get("blanketFrom"))))}</p><p>To: {escape(_line(_to_text(meta.get("blanketTo"))))}</p></div>
      <div class="panel"><h3>Certifier</h3><p>{escape(_line(_to_text(meta.get("certifierName"))))}</p><p>{escape(_line(_to_text(meta.get("certifierTitle"))))}</p><p>{escape(_line(_to_text(meta.get("certifierDate"))))}</p><p>{escape(_line(_to_text(meta.get("certifierContact"))))}</p></div>
    </div>
    """
    return _wrap_shipping_html("Certificate of Origin", body)


def _wrap_shipping_html(title: str, body: str) -> str:
    return f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>{escape(title)}</title>
  <style>
    :root {{
      --border: #d4d4d8;
      --muted: #666;
      --heading: #222;
    }}
    * {{ box-sizing: border-box; }}
    body {{
      margin: 24px;
      color: #111;
      font: 14px/1.45 "Segoe UI", Arial, sans-serif;
    }}
    h1 {{
      margin: 0;
      font-size: 24px;
      color: var(--heading);
    }}
    h3 {{
      margin: 0 0 8px;
      font-size: 12px;
      letter-spacing: .04em;
      text-transform: uppercase;
      color: #555;
    }}
    p {{ margin: 2px 0; }}
    .doc-header {{
      margin-bottom: 16px;
      padding-bottom: 10px;
      border-bottom: 1px solid var(--border);
    }}
    .doc-header p {{ color: var(--muted); }}
    .grid {{
      display: grid;
      gap: 12px;
      margin-bottom: 14px;
    }}
    .grid.two {{
      grid-template-columns: repeat(2, minmax(0, 1fr));
    }}
    .panel {{
      border: 1px solid var(--border);
      border-radius: 6px;
      padding: 10px;
    }}
    table {{
      width: 100%;
      border-collapse: collapse;
      margin-bottom: 14px;
    }}
    th, td {{
      border: 1px solid var(--border);
      padding: 8px;
      text-align: left;
      vertical-align: top;
    }}
    thead th {{
      background: #f6f6f7;
    }}
    .muted {{
      font-size: 12px;
      color: var(--muted);
    }}
    .totals {{
      max-width: 360px;
      margin-left: auto;
      border: 1px solid var(--border);
      border-radius: 6px;
      padding: 10px;
    }}
    .totals > div {{
      display: flex;
      justify-content: space-between;
      gap: 10px;
      margin: 4px 0;
    }}
    @media print {{
      body {{ margin: 12mm; }}
    }}
    @media (max-width: 800px) {{
      .grid.two {{ grid-template-columns: 1fr; }}
    }}
  </style>
</head>
<body>
{body}
</body>
</html>"""


def _html_lines(value: str) -> str:
    text = _to_text(value)
    if not text:
        return "-"
    return "<br />".join(escape(line) for line in text.splitlines() if line.strip()) or "-"


def _line(value: str) -> str:
    text = _to_text(value)
    return text if text else "-"


def _usd(amount: float) -> str:
    return f"${amount:,.2f}"


def _machine_rows(shipping_data: dict[str, Any]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for machine in shipping_data.get("machines", []):
        if not isinstance(machine, dict):
            continue
        machine_name = _to_text(machine.get("model")) or _to_text(machine.get("machineName"))
        machine_type = _to_text(machine.get("machineType") or machine.get("type"))
        quantity = _to_text(machine.get("qty") or machine.get("quantity") or "1")
        quantity_value = _to_float(quantity) if quantity else 1.0
        if quantity_value <= 0:
            quantity_value = 1.0
        unit_price = _to_float(machine.get("unitPrice"))
        line_total = unit_price * quantity_value
        rows.append(
            {
                "Machine": machine_name,
                "Serial_Number": _to_text(machine.get("serialNumber")),
                "HS": _to_text(machine.get("hsCode")),
                "TYPE": machine_type,
                "type": machine_type,
                "Qty": quantity,
                "Price": _usd(unit_price),
                "Total_Price": _usd(line_total),
                "Next Record": "",
            }
        )

    if rows:
        return rows

    return [
        {
            "Machine": "",
            "Serial_Number": "",
            "HS": "",
            "TYPE": "",
            "type": "",
            "Qty": "",
            "Price": "",
            "Total_Price": "",
            "Next Record": "",
        }
    ]


def _fill_repeating_machine_rows(document: Document, machine_rows: list[dict[str, str]]) -> None:
    required_count = max(1, len(machine_rows))
    for table in document.tables:
        row_indices = _find_repeating_row_indices(table)
        if not row_indices:
            continue

        if required_count > len(row_indices):
            template_row = table.rows[row_indices[-1]]._tr
            insert_before = table.rows[row_indices[-1] + 1]._tr if row_indices[-1] + 1 < len(table.rows) else None
            for _ in range(required_count - len(row_indices)):
                cloned = deepcopy(template_row)
                if insert_before is None:
                    table._tbl.append(cloned)
                else:
                    insert_before.addprevious(cloned)

            row_indices = _find_repeating_row_indices(table)

        for slot, row_index in enumerate(row_indices):
            row = table.rows[row_index]
            replacement = machine_rows[slot] if slot < len(machine_rows) else machine_rows[-1]
            if slot >= len(machine_rows):
                replacement = {
                    "Machine": "",
                    "Serial_Number": "",
                    "HS": "",
                    "TYPE": "",
                    "type": "",
                    "Qty": "",
                    "Price": "",
                    "Total_Price": "",
                    "Next Record": "",
                }
            for cell in row.cells:
                _replace_tokens_in_cell(cell, replacement, {})


def _find_repeating_row_indices(table: Any) -> list[int]:
    explicit_marker_rows = [
        index
        for index, row in enumerate(table.rows)
        if _row_contains_token(_table_row_text(row), "Next Record")
    ]
    if explicit_marker_rows:
        return explicit_marker_rows

    # Fallback for templates that omit Next Record but have machine placeholders.
    machine_template_tokens = ("Machine", "Serial_Number", "HS", "Qty", "TYPE", "type")
    fallback_rows: list[int] = []
    for index, row in enumerate(table.rows):
        row_text = _table_row_text(row)
        if any(_row_contains_token(row_text, token) for token in machine_template_tokens):
            fallback_rows.append(index)
    return fallback_rows


def _replace_tokens_in_document(
    document: Document,
    token_replacements: dict[str, str],
    literal_replacements: dict[str, str],
) -> None:
    _replace_tokens_in_container(document, token_replacements, literal_replacements)

    for section in document.sections:
        containers = (
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        )
        for container in containers:
            _replace_tokens_in_container(container, token_replacements, literal_replacements)


def _replace_tokens_in_container(
    container: Any,
    token_replacements: dict[str, str],
    literal_replacements: dict[str, str],
) -> None:
    for paragraph in container.paragraphs:
        _replace_tokens_in_paragraph(paragraph, token_replacements, literal_replacements)

    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                _replace_tokens_in_cell(cell, token_replacements, literal_replacements)


def _replace_tokens_in_cell(
    cell: Any,
    token_replacements: dict[str, str],
    literal_replacements: dict[str, str],
) -> None:
    for paragraph in cell.paragraphs:
        _replace_tokens_in_paragraph(paragraph, token_replacements, literal_replacements)
    for nested in cell.tables:
        for row in nested.rows:
            for nested_cell in row.cells:
                _replace_tokens_in_cell(nested_cell, token_replacements, literal_replacements)


def _replace_tokens_in_paragraph(
    paragraph: Any,
    token_replacements: dict[str, str],
    literal_replacements: dict[str, str],
) -> None:
    raw_text = "".join(run.text for run in paragraph.runs)
    if not raw_text:
        raw_text = paragraph.text or ""
    if not raw_text:
        return

    updated_text = raw_text
    for token, value in token_replacements.items():
        updated_text = updated_text.replace(f"\u00ab{token}\u00bb", value)
        updated_text = re.sub(r"\{\{\s*" + re.escape(token) + r"\s*\}\}", value, updated_text)

    for source, target in literal_replacements.items():
        if source:
            updated_text = updated_text.replace(source, target)

    updated_text = TOKEN_PATTERN.sub("", updated_text)
    updated_text = DOUBLE_BRACE_TOKEN_PATTERN.sub("", updated_text)
    updated_text = re.sub(r"\s+\n", "\n", updated_text)

    if updated_text == raw_text:
        return

    if paragraph.runs:
        paragraph.runs[0].text = updated_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = updated_text


def _build_global_token_replacements(
    shipping_data: dict[str, Any],
    machine_rows: list[dict[str, str]],
) -> dict[str, str]:
    client = shipping_data.get("client") if isinstance(shipping_data.get("client"), dict) else {}
    meta = shipping_data.get("meta") if isinstance(shipping_data.get("meta"), dict) else {}
    machines = [machine for machine in shipping_data.get("machines", []) if isinstance(machine, dict)]
    first_machine = machine_rows[0] if machine_rows else {"Machine": "", "Serial_Number": "", "HS": ""}
    display_name = _to_text(client.get("customerName")) or _to_text(client.get("company"))
    now = datetime.now()
    document_date = now.strftime("%Y-%m-%d")
    document_year = str(now.year)
    try:
        next_year_date = now.replace(year=now.year + 1).strftime("%Y-%m-%d")
    except ValueError:
        # Handle leap-day rollover gracefully for non-leap following years.
        next_year_date = now.replace(month=2, day=28, year=now.year + 1).strftime("%Y-%m-%d")
    manual_total = _to_float(meta.get("totalInvoiceAmount"))
    computed_total = sum(_to_float(machine.get("unitPrice")) for machine in machines)
    total_price = manual_total if manual_total > 0 else computed_total

    replacements = {
        "Company": display_name,
        "Customer": display_name,
        "Sold_toAddress_1": _to_text(client.get("soldToAddress1")),
        "Sold_toAddress_2": _to_text(client.get("soldToAddress2")),
        "Sold_toAddress_3": _to_text(client.get("soldToAddress3")),
        "Ship_toAddress_1": _to_text(client.get("shipToAddress1")),
        "Ship_toAddress_2": _to_text(client.get("shipToAddress2")),
        "Ship_toAddress_3": _to_text(client.get("shipToAddress3")),
        "Customer_PO": _to_text(client.get("customerPO")),
        "Order_date": _to_text(client.get("orderDate")),
        "Date": document_date,
        "Year": document_year,
        "Next_Year_Date": next_year_date,
        "Ox": _to_text(client.get("ox")),
        "Incoterm": _to_text(client.get("incoterm")),
        "Customer_Number": _to_text(client.get("customerNumber")),
        "Tax_ID": _to_text(client.get("taxId")),
        "Telefone": _to_text(client.get("telephone")),
        "Customer_contact": _to_text(client.get("clientContact")),
        "Machine": _to_text(first_machine.get("Machine")),
        "Serial_Number": _to_text(first_machine.get("Serial_Number")),
        "HS": _to_text(first_machine.get("HS")),
        "Qty": _to_text(first_machine.get("Qty")),
        "TYPE": _to_text(first_machine.get("TYPE")),
        "type": _to_text(first_machine.get("type")),
        "Price": _to_text(first_machine.get("Price")),
        "Total_Price": _usd(total_price),
        "Origin_Criterion": _to_text(meta.get("originCriterion") or "B"),
        "Country_of_origin": _to_text(meta.get("countryOfOrigin") or "Canada"),
        "Certifier_Name": _to_text(meta.get("certifierName")),
        "Certifier_Title": _to_text(meta.get("certifierTitle")),
        "Certifier_Date": _to_text(meta.get("certifierDate")),
    }
    return {key: value for key, value in replacements.items()}


def _build_literal_replacements(shipping_data: dict[str, Any]) -> dict[str, str]:
    machines = [machine for machine in shipping_data.get("machines", []) if isinstance(machine, dict)]
    trucks = [truck for truck in shipping_data.get("trucks", []) if isinstance(truck, dict)]
    meta = shipping_data.get("meta") if isinstance(shipping_data.get("meta"), dict) else {}

    total_crates = 0
    total_weight = 0.0
    for machine in machines:
        crates = machine.get("crates")
        if not isinstance(crates, list):
            continue
        total_crates += len(crates)
        for crate in crates:
            if not isinstance(crate, dict):
                continue
            total_weight += _to_float(crate.get("weightLbs"))

    truck_names = ", ".join(
        _to_text(truck.get("name")) for truck in trucks if _to_text(truck.get("name"))
    )

    replacements = {
        "X Crates (Total)": f"{total_crates} Crates (Total)",
        "x lbs (Total)": f"{total_weight:.1f} lbs (Total)",
        "Truck x": truck_names or "Truck 1",
    }

    broker_name = _to_text(meta.get("brokerInfo"))
    if broker_name:
        replacements["Reciever Broker"] = broker_name

    return replacements


def _build_truck_payloads(shipping_data: dict[str, Any]) -> list[dict[str, Any]]:
    machines = [machine for machine in shipping_data.get("machines", []) if isinstance(machine, dict)]
    trucks = [truck for truck in shipping_data.get("trucks", []) if isinstance(truck, dict)]

    if len(trucks) <= 1:
        return [{"index": 1, "shipping_data": dict(shipping_data)}]

    payloads: list[dict[str, Any]] = []
    for index, truck in enumerate(trucks, start=1):
        truck_id = _to_text(truck.get("id")) or f"truck-{index}"
        truck_name = _to_text(truck.get("name")) or f"Truck {index}"
        truck_machines = [
            machine
            for machine in machines
            if _to_text(machine.get("truckId")) == truck_id
        ]

        truck_state = dict(shipping_data)
        truck_state["machines"] = truck_machines
        truck_state["trucks"] = [{"id": truck_id, "name": truck_name}]
        payloads.append({"index": index, "shipping_data": truck_state})

    return payloads


def _build_line_item_options(line_item_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    options: list[dict[str, Any]] = []
    seen_keys: set[tuple[str, str, int]] = set()

    for index, line_item in enumerate(line_item_rows, start=1):
        if not isinstance(line_item, dict):
            continue

        description = _to_text(line_item.get("item_description"))
        if not description:
            continue
        name = _derive_line_item_name(description)

        quantity = _to_text(line_item.get("item_quantity"))
        unit_price = round(
            _to_float(line_item.get("item_price_numeric") or line_item.get("item_price_str")),
            2,
        )

        dedupe_key = (description.lower(), quantity, int(unit_price * 100))
        if dedupe_key in seen_keys:
            continue
        seen_keys.add(dedupe_key)

        raw_id = _to_int(line_item.get("id"))
        option_id = f"line-item-{raw_id if raw_id is not None else index}"
        options.append(
            {
                "id": option_id,
                "lineItemId": raw_id,
                "name": name or description,
                "description": description,
                "quantity": quantity,
                "unitPrice": unit_price,
            }
        )

    return options


def _resolve_machine_model(machine_name: str, machine_payload: dict[str, Any]) -> str:
    description = _to_text(machine_payload.get("description"))
    if description:
        return description.splitlines()[0].strip()

    main_item = machine_payload.get("main_item")
    if isinstance(main_item, dict):
        main_description = _to_text(main_item.get("description"))
        if main_description:
            return main_description.splitlines()[0].strip()

    return machine_name


def _derive_line_item_name(description: str) -> str:
    first_line = _normalize_machine_name_fragment(description.splitlines()[0] if description else "")
    if not first_line:
        return ""

    model_match = re.search(r"\bmodel\b[:\s-]*(.+)$", first_line, flags=re.IGNORECASE)
    if model_match:
        model_name = _normalize_machine_name_fragment(model_match.group(1))
        if model_name:
            return model_name

    return first_line


def _normalize_machine_name_fragment(value: str) -> str:
    if not value:
        return ""
    cleaned = re.sub(r"\s+", " ", value).strip(" \t-:;,.")
    # Normalize split model codes like "X65- 128" -> "X65-128"
    cleaned = re.sub(r"([A-Za-z0-9])\s*-\s*([A-Za-z0-9])", r"\1-\2", cleaned)
    return cleaned


def _resolve_machine_price(machine_payload: dict[str, Any]) -> float:
    total = 0.0

    main_item = machine_payload.get("main_item")
    if isinstance(main_item, dict):
        total += _to_float(main_item.get("item_price_numeric"))

    add_ons = machine_payload.get("add_ons")
    if isinstance(add_ons, list):
        for row in add_ons:
            if isinstance(row, dict):
                total += _to_float(row.get("item_price_numeric"))

    return round(total, 2)


def _split_address_lines(raw_value: Any) -> tuple[str, str, str]:
    text = _to_text(raw_value)
    if not text:
        return "", "", ""
    lines = [line.strip() for line in re.split(r"\r?\n", text) if line.strip()]
    return (
        lines[0] if len(lines) > 0 else "",
        lines[1] if len(lines) > 1 else "",
        "\n".join(lines[2:]) if len(lines) > 2 else "",
    )


def _table_row_text(row: Any) -> str:
    return " | ".join(cell.text for cell in row.cells)


def _row_contains_token(text: str, token_name: str) -> bool:
    if not text or not token_name:
        return False
    if f"\u00ab{token_name}\u00bb" in text:
        return True
    return bool(re.search(r"\{\{\s*" + re.escape(token_name) + r"\s*\}\}", text))


def _to_text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _to_float(value: Any) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    cleaned = re.sub(r"[^0-9.\-]", "", str(value))
    if not cleaned:
        return 0.0
    try:
        return float(cleaned)
    except ValueError:
        return 0.0


def _to_int(value: Any) -> int | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        return None


def _slugify(value: str) -> str:
    normalized = re.sub(r"[^a-zA-Z0-9]+", "-", value.strip()).strip("-")
    return normalized.lower()


def _normalize_machine_text(value: str) -> str:
    normalized = re.sub(r"[^a-z0-9]+", " ", value.lower())
    return re.sub(r"\s+", " ", normalized).strip()
